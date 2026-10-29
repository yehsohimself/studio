#!/usr/bin/env python3
"""
transkript.py - Bir video linkini indirir, konuşmayı yazıya döker, konuşmacıları ayırır.

Kullanım:
    python transkript.py <video_linki> [dil] [model] [seçenekler]

    dil   : auto | multi | tr | en | de | fr | es | ... (varsayılan: auto)
            auto  -> dili kendisi bulur; birden fazla dil sezerse karışık moda geçer
            multi -> karışık dil: her 30 saniyelik parçada dili yeniden tanır
    model : turbo | best        (varsayılan: turbo)
            turbo -> large-v3-turbo  (hızlı, günlük kullanım için yeterli)
            best  -> large-v3        (en doğru, 4-5 kat daha yavaş)

    Seçenekler:
      --isimler Ali,Ayşe       Konuşmacı adları (kim kim olduğu otomatik eşlenir)
      --sozluk onboarding,KPI   Sık geçen İngilizce terimler/özel adlar (ses modeline ipucu olur)
      --terim-duzelt            Yazı sonrası İngilizce terim düzeltmesi (deneysel; Türkçe kelimeleri de değiştirebilir)
      --ses-temizle             Çözümden önce gürültü bastırma ve ses dengeleme (gürültülü kayıtlar için)
      --konusmaci N             Konuşmacı sayısı belli ise ver (yoksa otomatik bulunur)
      --konusmaci-yok           Konuşmacı ayırma yapma (daha hızlı; tek konuşmacı olarak yazar)
      --altyazi-yok             Videoda hazır altyazı olsa bile kullanma, sesten çöz
      --ad DOSYA_ADI            Çıktı dosyalarının adı (uzantısız, varsayılan transkript_<dil>)
      --cerez chrome|safari     Giriş isteyen siteler (Instagram, TikTok) için tarayıcı çerezlerini kullan
      --goruntu                 Görüntüyü de incele: sahne kareleri, ekran yazıları, betimleme, senaryo çıktısı
      --ozet-yok                Özet çıkarma
      --claude-yok              Claude'u hiç kullanma (metin denetimi, ad önerisi, altyazı planı yok; Meet örnek toplantı testi)
      --ozet-dili tr|auto       Özetin dili (varsayılan tr; auto = videonun dili)

Çıktılar (bu klasöre, var olanın üzerine yazmaz):
    transkript_<dil>.txt   zaman damgalı + konuşmacı adlı metin, en altta düz metin
    transkript_<dil>.docx  aynı içerik Word belgesi olarak
    transkript_<dil>.json  ham veri (isimlendir.py ile konuşmacı adlarını sonradan değiştirmek için)
"""

# Instagram / TikTok gibi siteler giriş isterse bunu "chrome" veya "safari" yapın.
COOKIES_BROWSER = None
EK_TERIMLER = []          # --sozluk ile verilen ek terimler (ipucu + terim düzeltme)
SES_TEMIZLE = False       # --ses-temizle: çözümden önce gürültü bastırma (temizle.py)
TERIM_DUZELT = False      # --terim-duzelt ile açılır (gemma3:4b Türkçe kelimeleri de 'düzeltiyor'; varsayılan kapalı)

import itertools
import json


def _t(metin):
    """Belge başlıkları arayüz dilinde (dil.py; ayar `dil` == en ise İngilizce sözlük; 2026-10-25)."""
    try:
        import dil as _dil
        return _dil.t(metin)
    except Exception:  # noqa: BLE001
        return metin
import os
import re
import shutil
import subprocess
import sys
import threading
import tarfile
import tempfile
import time
import urllib.request
from pathlib import Path

MODELS = {"turbo": "large-v3-turbo", "best": "large-v3"}
MLX_MODELLER = {"large-v3-turbo": "mlx-community/whisper-large-v3-turbo", "large-v3": "mlx-community/whisper-large-v3-mlx"}


def motor_sec():
    """Ses çözme motoru: Apple grafik işlemcisi (MLX) varsa o, yoksa ana işlemci (faster-whisper)."""
    istenen = os.environ.get("TRANSKRIPT_MOTOR", "").lower()
    if istenen in ("faster", "cpu"):
        return "faster"
    try:
        import mlx.core  # noqa: F401
        import mlx_whisper  # noqa: F401
        return "mlx"
    except Exception:  # noqa: BLE001
        return "faster"
KLASOR = Path(__file__).resolve().parent
MODEL_KLASORU = KLASOR / "modeller"

SEG_URL = ("https://github.com/k2-fsa/sherpa-onnx/releases/download/"
           "speaker-segmentation-models/sherpa-onnx-pyannote-segmentation-3-0.tar.bz2")
SEG_MODEL = MODEL_KLASORU / "sherpa-onnx-pyannote-segmentation-3-0" / "model.onnx"
EMB_URL = ("https://github.com/k2-fsa/sherpa-onnx/releases/download/"
           "speaker-recongition-models/3dspeaker_speech_campplus_sv_zh_en_16k-common_advanced.onnx")
EMB_MODEL = MODEL_KLASORU / "3dspeaker_speech_campplus_sv_zh_en_16k-common_advanced.onnx"


# ----------------------------------------------------------------------------- yardımcılar
def zaman(saniye):
    saniye = int(saniye)
    return f"{saniye // 60:02d}:{saniye % 60:02d}"


def log(msg):
    print(msg, flush=True)


def ffmpeg_yolu():
    adaylar = [shutil.which("ffmpeg"), "/opt/homebrew/bin/ffmpeg", "/usr/local/bin/ffmpeg"]
    for aday in adaylar:
        if aday and Path(aday).exists():
            return str(Path(aday).parent)
    try:
        import imageio_ffmpeg
        return str(Path(imageio_ffmpeg.get_ffmpeg_exe()).parent)
    except Exception:
        pass
    try:
        import static_ffmpeg
        static_ffmpeg.add_paths()
        yol = shutil.which("ffmpeg")
        if yol:
            return str(Path(yol).parent)
    except Exception:
        pass
    return None


def modelleri_hazirla():
    """Konuşmacı ayırma modelleri yoksa indirir."""
    MODEL_KLASORU.mkdir(exist_ok=True)
    if not SEG_MODEL.exists():
        log("  Konuşmacı ayırma modeli indiriliyor (6 MB)...")
        tb = MODEL_KLASORU / "seg.tar.bz2"
        urllib.request.urlretrieve(SEG_URL, tb)
        with tarfile.open(tb, "r:bz2") as t:
            t.extractall(MODEL_KLASORU)
        tb.unlink()
    if not EMB_MODEL.exists():
        log("  Ses parmak izi modeli indiriliyor (28 MB)...")
        urllib.request.urlretrieve(EMB_URL, EMB_MODEL)


# ----------------------------------------------------------------------------- 1. indirme
def link_duzelt(link):
    """Bazı siteler doğrudan erişimi engeller; bilinen alternatif adresi kullan."""
    m = re.match(r"^https?://(?:www\.)?vimeo\.com/(?:video/)?(\d+)(?:/([0-9a-f]+))?", link)
    if m:
        yeni = f"https://player.vimeo.com/video/{m.group(1)}"
        if m.group(2):
            yeni += f"?h={m.group(2)}"
        return yeni
    return link


def yerel_dosya(link, hedef):
    """Link yerine bilgisayardaki bir ses/video dosyası verildiyse onu m4a'ya çevirir."""
    kaynak = Path(link).expanduser()
    ff = ffmpeg_yolu()
    log(f"[1/5] Dosya hazırlanıyor: {kaynak.name}")
    ses = Path(hedef) / "ses.m4a"
    komut = [str(Path(ff) / "ffmpeg") if ff else "ffmpeg", "-y", "-loglevel", "error", "-i", str(kaynak), "-vn", "-ac", "1", "-c:a", "aac", "-b:a", "96k", str(ses)]
    if subprocess.run(komut, env={**os.environ, "PATH": "/opt/homebrew/bin:/usr/local/bin:" + os.environ.get("PATH", "")}).returncode != 0 or not ses.exists():
        raise SystemExit("HATA: Dosya ses olarak okunamadı. Desteklenen bir video/ses dosyası mı?")
    sure = 0.0
    try:
        from faster_whisper.audio import decode_audio
        sure = len(decode_audio(str(ses), sampling_rate=16000)) / 16000
    except Exception:  # noqa: BLE001
        pass
    ad = re.sub(r"^[0-9a-f]{6}_", "", kaynak.stem)        # siteden yüklenen dosyaların geçici öneki
    ad = re.sub(r"\s*\[[^\]]*\]$", "", ad) or ad            # indirilenlerdeki "[720p]" gibi ekler
    return ses, None, {"baslik": ad, "kanal": "", "sure": sure, "kapak": "", "platform": "Dosya", "video_id": ""}


def indir(link, hedef, dil, altyazi_iste, video_iste=False):
    """Sesi m4a olarak ve varsa (yüklenmiş, otomatik olmayan) altyazıyı indirir.
    video_iste=True ise 480p videoyu da indirir (görüntü incelemesi için).
    Döner: (ses_yolu, altyazi_yolu|None, bilgi dict); bilgi["video"] video dosyası (varsa)"""
    hedef = Path(hedef)
    if not re.match(r"^https?://", link) and Path(link).expanduser().is_file():
        ses, alt, bilgi = yerel_dosya(link, hedef)
        bilgi["video"] = str(Path(link).expanduser()) if video_iste else ""
        return ses, alt, bilgi
    link = link_duzelt(link)
    komut = [sys.executable, "-m", "yt_dlp", "--no-playlist", "--no-simulate"]
    if video_iste:
        komut += ["-f", "bv*[height<=480][ext=mp4]+ba[ext=m4a]/bv*[height<=480]+ba/b[height<=480]/b",
                  "--merge-output-format", "mp4", "-o", str(hedef / "video.%(ext)s"),
                  "-k", "--extract-audio", "--audio-format", "m4a"]
    else:
        komut += ["-f", "bestaudio/best", "--extract-audio", "--audio-format", "m4a"]
    komut += [
        "-o", str(hedef / "ses.%(ext)s"),
        "--print-to-file", "%(title)s\n%(uploader)s\n%(duration)s\n%(thumbnail)s\n%(extractor_key)s\n%(id)s",
        str(hedef / "bilgi.txt"),
    ]
    if altyazi_iste:
        diller = "tr.*,en.*" if dil in ("auto", "multi") else f"{dil}.*,tr.*,en.*"
        komut += ["--write-subs", "--sub-langs", diller, "--sub-format", "vtt/srt/best"]
    ff = ffmpeg_yolu()
    if ff:
        komut += ["--ffmpeg-location", ff]
    if COOKIES_BROWSER:
        komut += ["--cookies-from-browser", COOKIES_BROWSER]

    env = dict(os.environ)
    env["PATH"] = "/opt/homebrew/bin:/usr/local/bin:" + env.get("PATH", "")
    deno = shutil.which("deno", path=env["PATH"])
    if deno:
        komut += ["--js-runtimes", f"deno:{deno}"]
    komut.append(link)

    log(f"[1/5] Ses indiriliyor: {link}")
    kayit = ""
    for deneme in range(1, 4):
        p = subprocess.Popen(komut, env=env, stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True)
        kayit = ""
        for satir in p.stdout:
            print(satir, end="", flush=True)
            kayit += satir
        p.wait()
        if p.returncode == 0:
            break
        k = kayit.lower()
        kalici = ("drm" in k or "login required" in k or "sign in" in k or "private video" in k
                  or "not available" in k or "404" in k or "unsupported url" in k)
        if kalici or deneme == 3:
            break
        log(f"  İndirme başarısız oldu, tekrar deneniyor ({deneme}/3)...")
        time.sleep(5)
    if p.returncode != 0:
        k = kayit.lower()
        if "drm" in k:
            raise SystemExit("HATA: Bu video kopya korumalı (DRM); ses indirilemiyor. "
                             "Bu tür videolar hiçbir araçla yazıya dökülemez.")
        if "login required" in k or "sign in" in k or "cookies" in k or "rate-limit" in k:
            raise SystemExit("HATA: Bu site giriş istiyor. Tarayıcı girişinizi kullanmak için "
                             "--cerez chrome (veya safari) ile tekrar deneyin.")
        if "private video" in k or "not available" in k or "404" in k:
            raise SystemExit("HATA: Video bulunamadı ya da gizli. Link doğru mu?")
        if "unsupported url" in k:
            raise SystemExit("HATA: Bu site desteklenmiyor ya da link bir video sayfası değil.")
        raise SystemExit("HATA: Ses indirilemedi. Link doğru mu? Giriş gerekiyorsa "
                         "--cerez chrome ile tekrar deneyin.")

    sesler = [p for p in hedef.glob("ses.*") if p.suffix in (".m4a", ".mp3", ".webm", ".opus", ".wav")]
    videolar = [p for p in hedef.glob("*") if p.suffix in (".mp4", ".mkv", ".webm") and p.stem != "ses"] + \
               [p for p in hedef.glob("ses.*") if p.suffix in (".mp4", ".mkv")]
    if not sesler and videolar:
        ses = hedef / "ses.m4a"
        subprocess.run([str(Path(ffmpeg_yolu() or "") / "ffmpeg") if ffmpeg_yolu() else "ffmpeg", "-y", "-loglevel", "error",
                        "-i", str(videolar[0]), "-vn", "-ac", "1", "-c:a", "aac", "-b:a", "96k", str(ses)], env=env)
        sesler = [ses] if ses.exists() else []
    if not sesler:
        raise SystemExit("HATA: İndirilen ses dosyası bulunamadı.")
    altyazilar = sorted(hedef.glob("ses.*.vtt")) + sorted(hedef.glob("ses.*.srt"))
    # Dil önceliği: istenen dil, sonra Türkçe, sonra İngilizce
    oncelik = ([dil] if dil != "auto" else []) + ["tr", "en"]
    altyazilar.sort(key=lambda p: (oncelik.index(altyazi_dili(p)) if altyazi_dili(p) in oncelik else 99))

    bilgi = {"baslik": "", "kanal": "", "sure": 0, "kapak": "", "platform": "", "video_id": ""}
    bt = hedef / "bilgi.txt"
    if bt.exists():
        satir = bt.read_text(encoding="utf-8").splitlines()
        al = lambda i: satir[i].strip() if len(satir) > i and satir[i].strip() != "NA" else ""  # noqa: E731
        bilgi["baslik"], bilgi["kanal"] = al(0), al(1)
        try:
            bilgi["sure"] = float(al(2) or 0)
        except Exception:
            pass
        bilgi["kapak"], bilgi["platform"], bilgi["video_id"] = al(3), al(4), al(5)
    bilgi["video"] = str(videolar[0]) if (video_iste and videolar) else ""
    return sesler[0], (altyazilar[0] if altyazilar else None), bilgi


# ----------------------------------------------------------------------------- 2. altyazı
def _vtt_zaman(s):
    s = s.strip().replace(",", ".")
    parca = s.split(":")
    parca = [float(p) for p in parca]
    while len(parca) < 3:
        parca.insert(0, 0.0)
    return parca[0] * 3600 + parca[1] * 60 + parca[2]


def altyazi_oku(yol):
    """VTT/SRT dosyasını (başlangıç, bitiş, metin) listesine çevirir."""
    metin = Path(yol).read_text(encoding="utf-8", errors="ignore")
    bloklar = re.split(r"\n\s*\n", metin)
    sonuc = []
    for blok in bloklar:
        satirlar = [s for s in blok.splitlines() if s.strip()]
        zi = next((i for i, s in enumerate(satirlar) if "-->" in s), None)
        if zi is None:
            continue
        bas, bit = satirlar[zi].split("-->")[:2]
        bit = bit.strip().split(" ")[0]
        try:
            b0, b1 = _vtt_zaman(bas), _vtt_zaman(bit)
        except Exception:
            continue
        govde = " ".join(satirlar[zi + 1:])
        govde = re.sub(r"<[^>]+>", "", govde)           # <c>, <i>, zaman etiketleri
        govde = re.sub(r"\s+", " ", govde).strip()
        if not govde:
            continue
        if sonuc and sonuc[-1][2] == govde and b0 <= sonuc[-1][1] + 0.05:
            sonuc[-1] = (sonuc[-1][0], b1, govde)       # yinelenen satır
            continue
        sonuc.append((b0, b1, govde))
    return sonuc


DIL_KODLARI = {"eng": "en", "tur": "tr", "deu": "de", "ger": "de", "fra": "fr", "fre": "fr",
               "spa": "es", "ita": "it", "rus": "ru", "ara": "ar", "por": "pt", "jpn": "ja",
               "kor": "ko", "zho": "zh", "chi": "zh", "nld": "nl", "dut": "nl"}


def altyazi_dili(yol):
    m = re.match(r"ses\.([A-Za-z\-_]+)\.(vtt|srt)$", Path(yol).name)
    if not m:
        return "auto"
    kod = re.split(r"[-_]", m.group(1))[0].lower()
    return DIL_KODLARI.get(kod, kod)


# ----------------------------------------------------------------------------- 3. ses → yazı
def _bellek_bosalt(model):
    """Ses modelini bellekten bırakır (görüntü ve özet modelleri için yer açar)."""
    import gc
    try:
        del model
    except Exception:  # noqa: BLE001
        pass
    gc.collect()


def dilleri_sez(dil_tani, ses_dosyasi, parcalar, pencere=12, ornek=None):
    """Videonun farklı yerlerinden örnek alıp hangi dillerin konuşulduğunu sezer.
    dil_tani(kesit) -> (dil, olasılık). Döner: [(dil, oran)] oran = pencerelerin payı, büyükten küçüğe."""
    import numpy as np
    from faster_whisper.audio import decode_audio
    if ornek is None:
        ornek = decode_audio(str(ses_dosyasi), sampling_rate=16000).astype(np.float32)
    toplam = len(ornek) / 16000
    pencere = max(pencere, min(24, int(toplam // 90)))      # uzun videoda daha çok örnek
    konusma = [(pb, pe) for pb, pe, _ in parcalar if pe - pb >= 1] or [(0.0, toplam)]

    def konusma_payi(b, e):
        return sum(max(0.0, min(e, pe) - max(b, pb)) for pb, pe in konusma)

    # Videonun tamamına eşit yayılmış 30 sn'lik pencereler; içinde yeterince konuşma olanları kullan
    adaylar = []
    for i in range(pencere):
        bas = max(0.0, min((i + 0.5) * toplam / pencere - 15, toplam - 30))
        if konusma_payi(bas, bas + 30) >= 8:
            adaylar.append(bas)
    sayac = {}
    for bas in sorted(set(round(a) for a in adaylar)):
        kesit = ornek[int(bas * 16000):int((bas + 30) * 16000)]
        if len(kesit) < 16000 * 5:
            continue
        try:
            d, p = dil_tani(kesit)
        except Exception:  # noqa: BLE001
            continue
        if p >= 0.6:
            sayac[d] = sayac.get(d, 0) + 1
    n = sum(sayac.values()) or 1
    return sorted(((d, c / n) for d, c in sayac.items()), key=lambda x: -x[1])


def transkript_et(ses_dosyasi, dil, model_adi, isimler, parcalar=None):
    if motor_sec() == "mlx":
        try:
            return transkript_et_mlx(ses_dosyasi, dil, model_adi, isimler, parcalar)
        except Exception as e:  # noqa: BLE001
            log(f"      Grafik işlemci motoru başarısız oldu ({e}); ana işlemciye geçiliyor.")
    return transkript_et_faster(ses_dosyasi, dil, model_adi, isimler, parcalar)


def _pencereler_olustur(konusma, toplam):
    """Konuşma aralıklarını ~12-40 sn'lik pencerelere böler. konusma: [(bas, bit)]"""
    if not konusma:
        konusma = [(float(i), min(i + 30.0, toplam)) for i in range(0, int(toplam), 30)]
    pencereler, pb0, pe0 = [], None, None
    for pb, pe in konusma:
        if pb0 is None:
            pb0, pe0 = pb, pe
        elif pb - pe0 > 1.2 and pe0 - pb0 >= 12 or pe - pb0 > 40:
            pencereler.append((pb0, pe0)); pb0, pe0 = pb, pe
        else:
            pe0 = max(pe0, pe)
    if pb0 is not None:
        pencereler.append((pb0, pe0))
    return pencereler


def _kosular(etiketli, bosluk=2.0):
    """Aynı dildeki ardışık pencereleri (aradaki boşluk küçükse) tek koşuda birleştirir."""
    kosular = []
    for pb, pe, d in etiketli:
        if kosular and kosular[-1][2] == d and pb - kosular[-1][1] <= bosluk:
            kosular[-1][1] = pe
        else:
            kosular.append([pb, pe, d])
    return kosular


HALUSINASYON_KALIPLARI = [
    "altyazı m.k", "altyazı: m.k", "izlediğiniz için teşekkür", "izlediğiniz için teşekkürler", "bir sonraki videoda görüşmek",
    "abone olmayı unutmayın", "beğenmeyi ve abone olmayı", "kanalıma abone", "thanks for watching", "thank you for watching",
    "subscribe to my channel", "see you in the next video", "please subscribe", "subtitles by", "amara.org",
]


def halusinasyon_mu(metin, dil, sn, toplam):
    """Whisper'ın sessizlikte uydurduğu bilindik kalıplar ve başka alfabeden çöp; video sonundaki gerçek vedalar korunur."""
    m = metin.lower().strip(" .!?,")
    if any(k in m for k in HALUSINASYON_KALIPLARI) and (toplam - sn > 90 or "altyazı" in m or "amara" in m):
        return True
    if dil in ("tr", "en", "de", "fr", "es", "it", "pt", "nl", "pl", "sv", "az", "id") and re.search(r"[\u0590-\u06ff\u0900-\u0dff\u0e00-\u0e7f\u3040-\u30ff\u4e00-\u9fff\uac00-\ud7af]", metin):
        return True
    if len(m) <= 3 and m in ("no", "ok", "the", "a", "i"):
        return True
    return False


def _yazisiz_araliklar(konusma, sonuc, en_az=5.0):
    """Konuşma (VAD) aralıklarının yazıya dökülmüş segmentlerle örtüşmeyen kısımlarını döner."""
    dolu = sorted((x[0], x[1]) for x in sonuc)
    eksik = []
    for pb, pe in konusma:
        t = pb
        for sb, se in dolu:
            if se <= t:
                continue
            if sb >= pe:
                break
            if sb - t >= en_az:
                eksik.append((t, sb))
            t = max(t, se)
        if pe - t >= en_az:
            eksik.append((t, pe))
    # Bitişikleri birleştir
    birlesik = []
    for a, b in eksik:
        if birlesik and a - birlesik[-1][1] <= 1.0:
            birlesik[-1] = (birlesik[-1][0], b)
        else:
            birlesik.append((a, b))
    return birlesik


def transkript_et_mlx(ses_dosyasi, dil, model_adi, isimler, parcalar=None):
    """Apple grafik işlemcisinde (MLX) çözüm. Sessizlikler atlanır; karışık dil pencere tabanlı."""
    import numpy as np
    import mlx.core as mx
    import mlx_whisper
    from mlx_whisper.transcribe import ModelHolder
    from mlx_whisper.audio import log_mel_spectrogram, pad_or_trim, N_FRAMES
    from faster_whisper.audio import decode_audio
    from faster_whisper.vad import VadOptions, get_speech_timestamps

    repo = MLX_MODELLER[model_adi]
    log(f"[3/5] Model yükleniyor: {model_adi} (grafik işlemci; ilk seferde indirilir)")
    model = ModelHolder.get_model(repo, mx.float16)
    ornek = decode_audio(str(ses_dosyasi), sampling_rate=16000).astype(np.float32)
    toplam = len(ornek) / 16000

    def dil_tani(kesit):
        mel = log_mel_spectrogram(kesit, n_mels=model.dims.n_mels)
        mel = pad_or_trim(mel, N_FRAMES, axis=-2).astype(mx.float16)
        _, olasiliklar = model.detect_language(mel)
        p = olasiliklar[0] if isinstance(olasiliklar, list) else olasiliklar
        d = max(p, key=p.get)
        return d, float(p[d])

    # Sessizlikleri atla: konuşma aralıklarını bul (dil sezimi de bunları kullanır)
    vad = get_speech_timestamps(ornek, VadOptions(min_silence_duration_ms=500, speech_pad_ms=300, threshold=0.4))
    konusma = [(v["start"] / 16000, v["end"] / 16000) for v in vad]
    coklu, diller = dil == "multi", []
    if dil in ("auto", "multi"):
        diller = dilleri_sez(dil_tani, ses_dosyasi, parcalar or [(a, b, 0) for a, b in konusma], ornek=ornek)
        if len([d for d, oran in diller if oran >= 0.15]) >= 2:
            coklu = True
        if diller:
            log("      Dil sezimi: " + ", ".join(f"{d} %{int(o * 100)}" for d, o in diller)
                + ("  → karışık dil modu" if coklu else ""))
    log("      Konuşma yazıya dökülüyor...")
    import terimler as terimler_modulu
    ipucu_yap = lambda d: terimler_modulu.ipucu_yap(d, isimler, EK_TERIMLER)   # noqa: E731

    pencereler = _pencereler_olustur(konusma, toplam)
    if coklu:
        varsayilan = diller[0][0] if diller else "tr"
        etiketli, son_dil = [], varsayilan
        for pb, pe in pencereler:
            kesit = ornek[int(max(0, pb - 0.2) * 16000):int(min(toplam, pe + 0.2) * 16000)]
            d, p = son_dil, 0.0
            if len(kesit) >= 16000 * 2:
                try:
                    d, p = dil_tani(kesit)
                except Exception:  # noqa: BLE001
                    pass
            if p < 0.5:
                d = son_dil
            etiketli.append((pb, pe, d)); son_dil = d
    else:
        sabit = None if dil == "auto" else dil
        if sabit is None:
            sabit = diller[0][0] if diller else None
        etiketli = [(pb, pe, sabit) for pb, pe in pencereler]
    kosular = _kosular(etiketli)
    if coklu:
        log(f"      {len(pencereler)} pencere, {len(kosular)} dil bloğu: " + ", ".join(f"{d} {zaman(pb)}-{zaman(pe)}" for pb, pe, d in kosular))

    sonuc, sayim = [], {}
    durum = {"onceki": ""}

    def coz(bas, bit, d, ipucu, kurtarma=False):
        """Bir aralığı çözer, elenen segmentleri atlayarak sonuc'a ekler. Döner: eklenen sayısı.
        kurtarma=True: sıcaklık 0, eşikler kapalı (model 'sessizlik' deyip atlamasın; eleme bizde)."""
        kesit = ornek[int(bas * 16000):int(bit * 16000)]
        if len(kesit) < 16000 * 0.5:
            return 0
        ek = {"temperature": 0.0, "compression_ratio_threshold": None, "logprob_threshold": None,
              "no_speech_threshold": None} if kurtarma else {"no_speech_threshold": 0.6, "temperature": (0.0, 0.2, 0.4)}
        r = mlx_whisper.transcribe(kesit, path_or_hf_repo=repo, language=d, word_timestamps=True,
                                   condition_on_previous_text=False, initial_prompt=ipucu, verbose=None, **ek)
        d_gercek = d or r.get("language") or "tr"
        sayim[d_gercek] = sayim.get(d_gercek, 0.0) + (bit - bas)
        eklenen = 0
        for seg in r.get("segments", []):
            metin = (seg.get("text") or "").strip()
            if not metin:
                continue
            if metin == durum["onceki"] and len(metin) > 15:
                continue
            if seg.get("avg_logprob", 0) < -1.2 and seg.get("no_speech_prob", 0) > 0.5:
                continue
            if seg.get("compression_ratio", 0) > 2.6:                 # takılma döngüsü
                continue
            if kurtarma and seg.get("avg_logprob", 0) < -1.5:          # kurtarmada gürültüyü alma
                continue
            if halusinasyon_mu(metin, d_gercek, float(seg["start"]) + bas, toplam):
                continue
            kelimeler = [[round(float(w["start"]) + bas, 2), round(float(w["end"]) + bas, 2), w["word"].strip()]
                         for w in (seg.get("words") or []) if w.get("word", "").strip()]
            guven = [round(float(w.get("probability", 1.0)), 2) for w in (seg.get("words") or []) if w.get("word", "").strip()]   # sözcük güveni (23 Ekim)
            sb = kelimeler[0][0] if kelimeler else float(seg["start"]) + bas
            se = kelimeler[-1][1] if kelimeler else float(seg["end"]) + bas
            if kurtarma and any(abs(x[0] - sb) < 1.0 and x[2] == metin for x in sonuc):   # zaten var
                continue
            sonuc.append((sb, se, metin, kelimeler, guven))
            durum["onceki"] = metin
            eklenen += 1
            log(f"  [{zaman(sb)}] {metin}")
        return eklenen

    for i, (pb, pe, d) in enumerate(kosular):
        coz(max(0.0, pb - 0.3), min(toplam, pe + 0.3), d, ipucu_yap(d or (diller[0][0] if diller else None)))

    # Kurtarma geçişi: konuşma var ama yazı çıkmamış aralıkları (≥5 sn) daha küçük parçalarla, ipuçsuz,
    # sabit sıcaklıkla yeniden çöz (takılma döngüsü ya da düşük güven yüzünden elenen pencereler için).
    eksikler = _yazisiz_araliklar(konusma, sonuc, en_az=5.0)
    if eksikler:
        log(f"      {len(eksikler)} aralıkta konuşma var ama yazı çıkmadı; yeniden deneniyor...")
        dil_bul = {}
        for pb, pe, d in kosular:
            dil_bul[(pb, pe)] = d
        kurtarilan = 0
        for eb, ee in eksikler:
            d = next((d for (pb, pe), d in dil_bul.items() if pb <= eb < pe), None) or (diller[0][0] if diller else None)
            adim = 12.0
            t = eb
            while t < ee:
                if min(t + adim, ee) - t >= 1.5:
                    kurtarilan += coz(max(0.0, t - 0.2), min(toplam, min(t + adim, ee) + 0.2), d, None, kurtarma=True)
                t += adim
        if kurtarilan:
            log(f"      {kurtarilan} parça kurtarıldı.")
        sonuc.sort(key=lambda x: x[0])
    sirali = sorted(sayim, key=lambda k: -sayim[k]) or ["tr"]
    transkript_et.diller = [d for d in sirali if sayim[d] >= 0.08 * toplam] or sirali[:1]
    # Belleği bırak
    ModelHolder.model = None; ModelHolder.model_path = None
    del model, ornek
    try:
        mx.clear_cache()
    except Exception:  # noqa: BLE001
        pass
    import gc; gc.collect()
    return transkript_et.diller[0], sonuc


def transkript_et_faster(ses_dosyasi, dil, model_adi, isimler, parcalar=None):
    from faster_whisper import WhisperModel

    log(f"[3/5] Model yükleniyor: {model_adi} (ana işlemci; ilk seferde indirilir, biraz sürebilir)")
    model = WhisperModel(model_adi, device="auto", compute_type="auto",
                         cpu_threads=os.cpu_count() or 4)
    coklu, diller = dil == "multi", []
    if dil in ("auto", "multi"):
        diller = dilleri_sez(lambda kesit: model.detect_language(audio=kesit)[:2], ses_dosyasi, parcalar or [])
        if len([d for d, oran in diller if oran >= 0.15]) >= 2:
            coklu = True
        if diller:
            log("      Dil sezimi: " + ", ".join(f"{d} %{int(o * 100)}" for d, o in diller)
                + ("  → karışık dil modu" if coklu else ""))
    log("      Konuşma yazıya dökülüyor...")
    import terimler as terimler_modulu
    ipucu_yap = lambda d: terimler_modulu.ipucu_yap(d, isimler, EK_TERIMLER)   # noqa: E731
    ayar = dict(beam_size=5, vad_filter=True,
                vad_parameters=dict(min_silence_duration_ms=500, speech_pad_ms=300, threshold=0.4),
                word_timestamps=True, condition_on_previous_text=False, no_speech_threshold=0.6)
    sonuc = []
    onceki = ""

    def topla(segments, kaydir=0.0):
        nonlocal onceki
        for seg in segments:
            metin = seg.text.strip()
            if not metin:
                continue
            if metin == onceki and len(metin) > 15:          # takılma / tekrar döngüsü koruması
                continue
            if seg.avg_logprob is not None and seg.avg_logprob < -1.2 and seg.no_speech_prob > 0.5:
                continue
            bas, bit = seg.start, seg.end
            kelimeler = []
            if seg.words:
                bas, bit = seg.words[0].start, seg.words[-1].end
                kelimeler = [[round(w.start + kaydir, 2), round(w.end + kaydir, 2), w.word.strip()] for w in seg.words if w.word.strip()]
            guven = [round(float(getattr(w, "probability", 1.0) or 1.0), 2) for w in (seg.words or []) if w.word.strip()]
            sonuc.append((bas + kaydir, bit + kaydir, metin, kelimeler, guven))
            onceki = metin
            log(f"  [{zaman(bas + kaydir)}] {metin}")

    if not coklu:
        tek_dil = None if dil == "auto" else dil
        segments, info = model.transcribe(str(ses_dosyasi), language=tek_dil,
                                          initial_prompt=ipucu_yap(tek_dil or (diller[0][0] if diller else None)), **ayar)
        topla(segments)
        transkript_et.diller = [info.language]
        _bellek_bosalt(model)
        return info.language, sonuc

    # Karışık dil: konuşma aralıklarını pencerelere böl, her pencerenin dilini tanı,
    # aynı dildeki ardışık pencereleri o dil sabitlenmiş olarak çöz.
    import numpy as np
    from faster_whisper.audio import decode_audio
    ornek = decode_audio(str(ses_dosyasi), sampling_rate=16000).astype(np.float32)
    toplam = len(ornek) / 16000
    konusma = sorted((pb, pe) for pb, pe, _ in (parcalar or []) if pe > pb)
    if not konusma:
        konusma = [(i, min(i + 30.0, toplam)) for i in range(0, int(toplam), 30)]
    pencereler, pb0, pe0 = [], None, None
    for pb, pe in konusma:
        if pb0 is None:
            pb0, pe0 = pb, pe
        elif pb - pe0 > 1.2 and pe0 - pb0 >= 12 or pe - pb0 > 40:
            pencereler.append((pb0, pe0)); pb0, pe0 = pb, pe
        else:
            pe0 = max(pe0, pe)
    if pb0 is not None:
        pencereler.append((pb0, pe0))
    varsayilan = diller[0][0] if diller else "tr"
    etiketli, son_dil = [], varsayilan
    for pb, pe in pencereler:
        kesit = ornek[int(max(0, pb - 0.2) * 16000):int(min(toplam, pe + 0.2) * 16000)]
        d, p = son_dil, 0.0
        if len(kesit) >= 16000 * 2:
            try:
                d, p, _ = model.detect_language(audio=kesit)
            except Exception:  # noqa: BLE001
                pass
        if p < 0.5:
            d = son_dil
        etiketli.append([pb, pe, d]); son_dil = d
    # Aynı dildeki ardışık pencereleri birleştir
    kosular = []
    for pb, pe, d in etiketli:
        if kosular and kosular[-1][2] == d:
            kosular[-1][1] = pe
        else:
            kosular.append([pb, pe, d])
    log(f"      {len(pencereler)} pencere, {len(kosular)} dil bloğu: " + ", ".join(f"{d} {zaman(pb)}-{zaman(pe)}" for pb, pe, d in kosular))
    sayim = {}
    for pb, pe, d in kosular:
        sayim[d] = sayim.get(d, 0.0) + (pe - pb)
    for i, (pb, pe, d) in enumerate(kosular):
        bas = 0.0 if i == 0 else max(0.0, pb - 0.3)
        bit = toplam if i == len(kosular) - 1 else min(toplam, pe + 0.3)
        kesit = ornek[int(bas * 16000):int(bit * 16000)]
        if len(kesit) < 16000:
            continue
        segments, _ = model.transcribe(kesit, language=d, initial_prompt=ipucu_yap(d), **ayar)
        topla(segments, kaydir=bas)
    sirali = sorted(sayim, key=lambda k: -sayim[k])
    transkript_et.diller = [d for d in sirali if sayim[d] >= 0.08 * toplam] or sirali[:1]
    del ornek
    _bellek_bosalt(model)
    return transkript_et.diller[0], sonuc


# ----------------------------------------------------------------------------- 4. konuşmacılar
def _performans_cekirdegi():
    """Apple Silicon'da performans çekirdeği sayısı (yoksa en çok 4)."""
    try:
        import subprocess
        n = int(subprocess.run(["sysctl", "-n", "hw.perflevel0.physicalcpu"], capture_output=True, text=True, timeout=3).stdout.strip())
        return max(1, min(n, 8))
    except Exception:  # noqa: BLE001
        return min(4, max(1, (os.cpu_count() or 4) - 1))


def _ayirma_baslat(ses, konusmaci_sayisi, cikti):
    """Konuşmacı ayırmayı ayrı bir süreçte başlatır (sonuç JSON dosyasına yazılır)."""
    import subprocess
    log("[2/5] Konuşmacılar ayrıştırılıyor (yazıya dökmeyle aynı anda)...")
    kod = ("import sys, json, numpy as np; sys.path.insert(0, %r); import transkript as T; T.log = lambda *a, **k: None; "
           "p = T.konusmacilari_ayir(sys.argv[1], int(sys.argv[2]) if sys.argv[2] != '0' else None); "
           "iz = {str(k): [float(x) for x in np.asarray(v).ravel()] for k, v in (getattr(T.konusmacilari_ayir, 'izler', {}) or {}).items()}; "
           "json.dump({'parcalar': p, 'izler': iz}, open(sys.argv[3], 'w'))") % str(KLASOR)
    return subprocess.Popen([sys.executable, "-c", kod, str(ses), str(konusmaci_sayisi or 0), str(cikti)],
                            stdout=subprocess.DEVNULL, stderr=subprocess.PIPE, text=True), cikti


def _ayirma_bitir(ayirici, ses, konusmaci_sayisi):
    """Ayrı süreçteki ayırmayı bekler; başarısızsa aynı süreçte yeniden dener."""
    proc, cikti = ayirici
    _out, err = proc.communicate()
    try:
        d = json.loads(Path(cikti).read_text(encoding="utf-8"))
        parcalar = [(float(a), float(b), int(k)) for a, b, k in d["parcalar"]]
        konusmacilari_ayir.izler = {k: [float(x) for x in v] for k, v in d.get("izler", {}).items()}   # liste kalsın (JSON/ogren)
        log(f"      {len({p[2] for p in parcalar})} konuşmacı bulundu.")
        return parcalar
    except Exception as e:  # noqa: BLE001
        log(f"      Paralel ayırma sonuç vermedi ({e}); yeniden deneniyor." + (f" [{err.strip()[-200:]}]" if err else ""))
        return konusmacilari_ayir(ses, konusmaci_sayisi)


def konusmacilari_ayir(ses_dosyasi, konusmaci_sayisi):
    """Kim ne zaman konuşuyor? -> [(bas, bit, konusmaci_no)]"""
    import numpy as np
    import sherpa_onnx
    from faster_whisper.audio import decode_audio

    modelleri_hazirla()
    log("[2/5] Konuşmacılar ayrıştırılıyor...")
    ornekler = decode_audio(str(ses_dosyasi), sampling_rate=16000).astype(np.float32)
    # Ölçüm (58 dk toplantı, M4): 9 iş parçacığı 714 sn, 4 iş parçacığı 205 sn (3,5x). Fazla iş parçacığı
    # verimlilik çekirdeklerine taşıp yavaşlatıyor; performans çekirdeği sayısı (4) en iyisi.
    nt = int(os.environ.get("SHERPA_NT") or 0) or _performans_cekirdegi()
    config = sherpa_onnx.OfflineSpeakerDiarizationConfig(
        segmentation=sherpa_onnx.OfflineSpeakerSegmentationModelConfig(
            pyannote=sherpa_onnx.OfflineSpeakerSegmentationPyannoteModelConfig(model=str(SEG_MODEL)),
            num_threads=nt,
        ),
        embedding=sherpa_onnx.SpeakerEmbeddingExtractorConfig(model=str(EMB_MODEL), num_threads=nt),
        clustering=sherpa_onnx.FastClusteringConfig(
            num_clusters=konusmaci_sayisi if konusmaci_sayisi else -1, threshold=0.9),
        min_duration_on=0.3,
        min_duration_off=0.5,
    )
    sd = sherpa_onnx.OfflineSpeakerDiarization(config)
    sonuc = sd.process(ornekler).sort_by_start_time()
    parcalar = [(s.start, s.end, int(s.speaker)) for s in sonuc]
    parcalar = kucuk_kumeleri_birlestir(parcalar)
    kisiler = sorted({p[2] for p in parcalar})
    log(f"      {len(kisiler)} konuşmacı bulundu.")
    try:
        import sesler as sesler_modulu
        konusmacilari_ayir.izler = sesler_modulu.izleri_cikar(ornekler, parcalar, EMB_MODEL, nt, log)
    except Exception as e:  # noqa: BLE001
        log(f"      Ses izi çıkarılamadı: {e}")
        konusmacilari_ayir.izler = {}
    return parcalar


def kucuk_kumeleri_birlestir(parcalar, oran=0.015, en_az=10.0):
    """Toplamda çok az konuşan 'hayalet' konuşmacıları komşu konuşmacıya katar."""
    if not parcalar:
        return parcalar
    toplam = sum(pe - pb for pb, pe, _ in parcalar)
    sure = {}
    for pb, pe, k in parcalar:
        sure[k] = sure.get(k, 0.0) + (pe - pb)
    buyukler = {k for k, t in sure.items() if t >= max(en_az, toplam * oran)}
    if not buyukler:
        buyukler = {max(sure, key=sure.get)}
    yeni = []
    for i, (pb, pe, k) in enumerate(parcalar):
        if k not in buyukler:
            # En yakın (zamanca) büyük konuşmacı parçasını bul
            en_yakin, mesafe = None, None
            for qb, qe, qk in parcalar:
                if qk in buyukler:
                    d = max(0.0, max(pb, qb) - min(pe, qe))
                    if mesafe is None or d < mesafe:
                        en_yakin, mesafe = qk, d
            k = en_yakin
        yeni.append((pb, pe, k))
    # Numaraları 0..n-1 yap
    sira = {k: i for i, k in enumerate(sorted(buyukler))}
    return [(pb, pe, sira[k]) for pb, pe, k in yeni]


def konusmaci_eslestir(metinler, parcalar):
    """Her metin parçasına en çok çakışan konuşmacıyı verir."""
    sonuc = []
    for satir in metinler:
        bas, bit, metin = satir[0], satir[1], satir[2]
        kelimeler = satir[3] if len(satir) > 3 else []
        puan = {}
        for pb, pe, k in parcalar:
            cak = min(bit, pe) - max(bas, pb)
            if cak > 0:
                puan[k] = puan.get(k, 0) + cak
        belirsiz = False
        if puan:
            k = max(puan, key=puan.get)
            # 23 Ekim (ChatGPT brifi, Meet): konuşmacı ataması belirsizse işaretlenir (ses başka bir konuşmacının parçasıyla da büyük ölçüde çakışıyor)
            belirsiz = len(puan) > 1 and puan[k] / sum(puan.values()) < 0.65
        else:
            # Çakışma yok: ortasına en yakın parça
            orta = (bas + bit) / 2
            k = min(parcalar, key=lambda p: min(abs(p[0] - orta), abs(p[1] - orta)))[2] if parcalar else 0
        d = {"bas": bas, "bit": bit, "k": k, "metin": metin}
        if kelimeler:
            d["kelimeler"] = kelimeler
            if len(satir) > 4 and satir[4] and len(satir[4]) == len(kelimeler):
                d["guven"] = satir[4]
        if belirsiz:
            d["k_belirsiz"] = True
        sonuc.append(d)
    return sonuc


def hayalet_temizle(segmentler, pay=0.06, medyan_sinir=2.5, zorla=None):
    """Toplamda az konuşan ve hep kısa parçalardan oluşan 'hayalet' konuşmacıyı siler;
    parçalarını zamanca komşu konuşmacıya verir. zorla=["2"] ile belli biri de silinebilir."""
    import statistics
    if not segmentler:
        return segmentler
    toplam = sum(s["bit"] - s["bas"] for s in segmentler) or 1.0
    kisiler = sorted({s["k"] for s in segmentler})
    if len(kisiler) < 2:
        return segmentler
    hayaletler = set(zorla or [])
    for k in kisiler:
        parca = [s["bit"] - s["bas"] for s in segmentler if s["k"] == k]
        if sum(parca) / toplam < pay and statistics.median(parca) < medyan_sinir:
            hayaletler.add(k)
    if len(hayaletler) >= len(kisiler):            # herkes hayalet olamaz
        hayaletler = set(zorla or [])
    if not hayaletler:
        return segmentler
    n = len(segmentler)
    for i, s in enumerate(segmentler):
        if s["k"] not in hayaletler:
            continue
        onceki = next((segmentler[j] for j in range(i - 1, -1, -1) if segmentler[j]["k"] not in hayaletler), None)
        sonraki = next((segmentler[j] for j in range(i + 1, n) if segmentler[j]["k"] not in hayaletler), None)
        if onceki and sonraki:
            s["k"] = onceki["k"] if (s["bas"] - onceki["bit"]) <= (sonraki["bas"] - s["bit"]) else sonraki["k"]
        else:
            s["k"] = (onceki or sonraki)["k"]
    return segmentler


def isimleri_esle(segmentler, isimler):
    """Verilen adları konuşmacı numaralarına dağıtır: insan kendi adını en az söyler."""
    kisiler = sorted({s["k"] for s in segmentler})
    etiket = {k: f"Konuşmacı {i + 1}" for i, k in enumerate(kisiler)}
    if not isimler:
        return etiket
    # kişi -> ad -> kaç kez söyledi
    sayac = {k: {ad: 0 for ad in isimler} for k in kisiler}
    for s in segmentler:
        for ad in isimler:
            sayac[s["k"]][ad] += len(re.findall(rf"\b{re.escape(ad)}\b", s["metin"], flags=re.IGNORECASE))
    # Konuşma süresi (uzun konuşanlara ad ver, kalanlar numara alır)
    sure = {k: 0.0 for k in kisiler}
    for s in segmentler:
        sure[s["k"]] += s["bit"] - s["bas"]
    adaylar = sorted(kisiler, key=lambda k: -sure[k])[:len(isimler)]
    en_iyi, en_iyi_puan = None, None
    for perm in itertools.permutations(isimler, len(adaylar)):
        puan = sum(sayac[k][ad] for k, ad in zip(adaylar, perm))   # kendi adını söyleme sayısı
        if en_iyi_puan is None or puan < en_iyi_puan:
            en_iyi, en_iyi_puan = perm, puan
    for k, ad in zip(adaylar, en_iyi):
        etiket[k] = ad
    return etiket


# ----------------------------------------------------------------------------- 5. birleştir / kontrol
def paragrafla(segmentler, etiket, max_bosluk=2.0, max_uzunluk=600):
    """Aynı kişinin ardışık cümlelerini paragraf yapar."""
    paragraflar = []
    for i, s in enumerate(segmentler):
        ad = etiket[s["k"]]
        if (paragraflar and paragraflar[-1]["ad"] == ad
                and s["bas"] - paragraflar[-1]["bit"] <= max_bosluk
                and len(paragraflar[-1]["metin"]) + len(s["metin"]) <= max_uzunluk):
            paragraflar[-1]["metin"] += " " + s["metin"]
            paragraflar[-1]["bit"] = s["bit"]
            paragraflar[-1]["si"].append(i)
        else:
            paragraflar.append({"bas": s["bas"], "bit": s["bit"], "ad": ad, "metin": s["metin"], "si": [i]})
    return paragraflar


def eksik_kontrol(segmentler, parcalar, esik=6.0):
    """Konuşma var ama yazı yok görünen yerleri bulur (olası kaçan metin)."""
    uyarilar = []
    for pb, pe, _ in parcalar:
        if pe - pb < esik:
            continue
        kapali = 0.0
        for s in segmentler:
            kapali += max(0.0, min(pe, s["bit"]) - max(pb, s["bas"]))
        if kapali < (pe - pb) * 0.3:
            uyarilar.append(f"[{zaman(pb)}-{zaman(pe)}] konuşma var ama yazıya dökülen çok az")
    return uyarilar


# ----------------------------------------------------------------------------- 6. çıktılar
def cikti_adi(temel):
    aday = KLASOR / temel
    n = 2
    while any((KLASOR / f"{aday.name}{u}").exists() for u in (".txt", ".docx", ".json")):
        aday = KLASOR / f"{temel}_{n}"
        n += 1
    return aday


def yaz_txt(yol, veri, paragraflar):
    with open(yol, "w", encoding="utf-8") as f:
        f.write(f"Başlık: {veri['baslik']}\n")
        f.write(f"Kanal: {veri['kanal']}\n")
        f.write(f"Link: {veri['link']}\n")
        f.write("Tespit edilen dil: " + ", ".join(veri.get("diller") or [veri["dil"]]) + "\n")
        f.write(f"{_t('Kaynak')}: {veri['kaynak']}\n")
        f.write(_t("Konuşmacılar") + ": " + ", ".join(veri["etiket"].values()) + "\n")
        if veri["uyarilar"]:
            f.write("\nKontrol notları:\n")
            for u in veri["uyarilar"]:
                f.write(f"  - {u}\n")
        o = veri.get("ozet")
        if o:
            f.write("\n--- Özet ---\n\n" + o["ozet"] + "\n")
            if o.get("maddeler"):
                f.write("\nÖne çıkanlar:\n")
                for m in o["maddeler"]:
                    f.write(f"  • {m}\n")
            if o.get("kararlar"):
                f.write("\nKararlar:\n")
                for m in o["kararlar"]:
                    f.write(f"  • {m}\n")
            if o.get("yapilacaklar"):
                f.write("\nYapılacaklar:\n")
                for m in o["yapilacaklar"]:
                    f.write(f"  ☐ {m}\n")
            if o.get("bolumler"):
                f.write("\nBölümler:\n")
                for b in o["bolumler"]:
                    f.write(f"  [{zaman_uzun(b['sn'], any(bb['sn'] >= 3600 for bb in o['bolumler']))}] {b['baslik']}\n")
        ks = veri.get("kisaca")
        if ks:
            f.write("\n--- Kısaca ---\n\n" + "\n\n".join(ks.get("ozet") or []) + "\n")
            for k in (ks.get("kisiler") or []) + ([{"ad": "Sorumlu belirtilmedi", "isler": ks["belirsiz"]}] if ks.get("belirsiz") else []):
                f.write(f"\n{k['ad']}:\n" + "".join(f"  {'☑' if x.get('bitti') else '☐'} {x['is']}" + (f" ({x['ne_zaman']})" if x.get("ne_zaman") else "") + (f" [{x['musteri']}]" if x.get("musteri") else "") + "\n" for x in k["isler"]))
        if ks and ks.get("kararlar"):
            f.write("\nKararlar:\n" + "".join(f"  • {x['karar']}" + (f" [{x['musteri']}]" if x.get("musteri") else "") + "\n" + "".join(f"    ↳ {_zincir_metni(z)}\n" for z in _zincirler(x)) for x in ks["kararlar"]))
        if ks and ks.get("oneriler"):
            f.write("\nÖneriler (karar değil):\n" + "".join(f"  – {x['oneri']}" + (f" [{x['musteri']}]" if x.get("musteri") else "") + "\n" for x in ks["oneriler"]))
        if ks and ks.get("acik_sorular"):
            f.write("\nAçık sorular:\n" + "".join(f"  ? {x['soru']}" + (f" [{x['musteri']}]" if x.get("musteri") else "") + "\n" for x in ks["acik_sorular"]))
        if veri.get("goruntu"):
            f.write("\n--- Görüntü notları ---\n\n")
            sg = any(g["sn"] >= 3600 for g in veri["goruntu"])
            for g in veri["goruntu"]:
                satir = f"[{zaman_uzun(g['sn'], sg)}] " + (g.get("aciklama") or "").strip()
                if g.get("yazi"):
                    satir += ("  " if g.get("aciklama") else "") + "Ekranda: " + " / ".join(g["yazi"])
                f.write(satir.strip() + "\n")
        f.write("\n--- Zaman damgalı ---\n\n")
        saatli = any(p["bas"] >= 3600 for p in paragraflar)
        for p in paragraflar:
            f.write(f"[{zaman_uzun(p['bas'], saatli)}] {p['ad']}: {p['metin']}\n\n")
        f.write("--- Düz metin ---\n\n")
        f.write(" ".join(p["metin"] for p in paragraflar) + "\n")


def _alan_ekle(run, alan):
    """Word alanı (ör. PAGE) ekler."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin")
    t = OxmlElement("w:instrText"); t.set(qn("xml:space"), "preserve"); t.text = alan
    e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end")
    run._r.append(b); run._r.append(t); run._r.append(e)


def sure_metni(saniye):
    """İnsan okur biçim: 1 sa 9 dk / 12 dk 5 sn / 52 sn."""
    sn = int(saniye or 0)
    h, m, s = sn // 3600, sn % 3600 // 60, sn % 60
    if h:
        return f"{h} sa {m} dk"
    return f"{m} dk {s} sn" if m else f"{s} sn"


def zaman_uzun(saniye, saatli=False):
    """Zaman damgası: mm:ss, 1 saati aşan videolarda h:mm:ss."""
    sn = int(saniye)
    if saatli:
        return f"{sn // 3600}:{sn % 3600 // 60:02d}:{sn % 60:02d}"
    return zaman(sn)


def _surum():
    """Video Hub'ın o anki sürümü (yerel git: kısa kod + tarih); kaydın künyesine yazılır."""
    try:
        p = subprocess.run(["git", "-C", str(Path(__file__).resolve().parent), "log", "-1", "--format=%h %cs"], capture_output=True, text=True,
                           timeout=5).stdout.split()
        return {"kod": p[0], "tarih": p[1]} if len(p) == 2 else None
    except Exception:  # noqa: BLE001
        return None


def _zincirler(x):
    return [dict(z, yon="once") for z in x.get("zincir") or []] + [dict(z, yon="sonra") for z in x.get("sonra") or []]


def _zincir_metni(z):
    """Karar zinciri satırı (karar_zinciri.py): '14 Ekim kararını değiştirdi: ...' ya da '23 Ekim'de değişti: ...'."""
    try:
        g = time.strftime("%d.%m", time.strptime(z.get("tarih") or "", "%Y-%m-%d"))
    except ValueError:
        g = z.get("tarih") or ""
    ac = z.get("aciklama") or z.get("karar") or ""
    if z.get("yon") == "once":
        return f"{g} kararını {({'degistirdi': 'değiştirdi', 'geri_aldi': 'geri aldı', 'teyit': 'teyit etti'}).get(z.get('iliski'), '')}: {ac}"
    return f"{g} toplantısında {({'degistirdi': 'değişti', 'geri_aldi': 'geri alındı', 'teyit': 'yeniden onaylandı'}).get(z.get('iliski'), '')}: {ac}"


def yaz_docx(yol, veri, paragraflar, yalniz_danisman=False):
    """yalniz_danisman: yalnızca kişisel danışman notu (ayrı dosya, Drive'a gitmez; 2026-10-23). Ana belgede danışman notu YOK (paylaşılırsa sızmasın)."""
    if yalniz_danisman:
        veri = {**{k: veri.get(k) for k in ("baslik", "danisman", "olusturma", "toplanti", "dil", "musteri") if veri.get(k) is not None},
                "dil": veri.get("dil") or "tr", "etiket": {}}
        paragraflar = []
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    GRI = RGBColor(0x6E, 0x6E, 0x73)
    KOYU = RGBColor(0x1D, 0x1D, 0x1F)
    RENKLER = [RGBColor(0x00, 0x66, 0xCC), RGBColor(0xC1, 0x44, 0x0E), RGBColor(0x1B, 0x7F, 0x45),
               RGBColor(0x7C, 0x3A, 0xED), RGBColor(0x0A, 0x7C, 0x8A), RGBColor(0x8A, 0x5A, 0x2B)]
    YAZI = "Helvetica Neue"
    saatli = any(pp["bas"] >= 3600 for pp in paragraflar)

    def font_ayarla(run, boyut, kalin=False, renk=None, tabular=False):
        run.font.name = YAZI
        run.font.size = Pt(boyut)
        run.bold = kalin
        if renk is not None:
            run.font.color.rgb = renk
        rpr = run._r.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts"); rpr.append(rf)
        for k in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rf.set(qn(k), YAZI)
        if tabular:
            ns = "http://schemas.microsoft.com/office/word/2010/wordml"
            el = OxmlElement("w14:numSpacing"); el.set("{%s}val" % ns, "tabular"); rpr.append(el)

    doc = Document()
    for sec in doc.sections:
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)        # A4
        sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.2)
        sec.left_margin, sec.right_margin = Cm(2.4), Cm(2.4)
    normal = doc.styles["Normal"]
    normal.font.name = YAZI
    normal.font.size = Pt(11)
    normal.font.color.rgb = KOYU
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.2
    rpr = normal.element.get_or_add_rPr()
    lang = OxmlElement("w:lang")
    for k in ("w:val", "w:eastAsia"):
        lang.set(qn(k), "tr-TR")
    rpr.append(lang)
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.append(rf)
    for k in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(k), YAZI)
    cp = doc.core_properties
    cp.title = veri.get("baslik") or "Transkript"
    cp.author = "Transkript"
    cp.subject = veri.get("kanal") or ""
    cp.language = "tr-TR"
    cp.comments = ""
    cp.keywords = "transkript"

    def par(metin="", boyut=11, kalin=False, renk=None, once=0, sonra=0, hiza=None):
        p = doc.add_paragraph()
        p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(once), Pt(sonra)
        if hiza is not None:
            p.alignment = hiza
        if metin:
            font_ayarla(p.add_run(metin), boyut, kalin, renk)
        return p

    # Başlık bloğu
    p = par("", sonra=6)
    r = p.add_run("TRANSKRİPT"); font_ayarla(r, 9, True, GRI)
    sp = OxmlElement("w:spacing"); sp.set(qn("w:val"), "30"); r._r.get_or_add_rPr().append(sp)
    par(veri["baslik"] or "Transkript", 24, True, KOYU, sonra=4)
    try:
        from ozet import DIL_ADI
    except Exception:  # noqa: BLE001
        DIL_ADI = {}
    dil_adi = " + ".join(DIL_ADI.get(d, d) for d in (veri.get("diller") or [veri["dil"]]))
    kaynak = "Hazır altyazı" if "altyazı" in (veri.get("kaynak") or "") else "Sesten çözüldü"
    bilgi = [veri.get("kanal"), dil_adi, sure_metni(veri["sure"]) if veri.get("sure") else None, kaynak]
    par("  ·  ".join(b for b in bilgi if b), 10, False, GRI, sonra=3)

    # Köprü: Videoyu aç
    if veri.get("link"):
        p = par("", sonra=4)
        rel_id = doc.part.relate_to(veri["link"], "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
        h = OxmlElement("w:hyperlink"); h.set(qn("r:id"), rel_id)
        nr = OxmlElement("w:r"); rp = OxmlElement("w:rPr")
        c = OxmlElement("w:color"); c.set(qn("w:val"), "0066CC"); rp.append(c)
        sz = OxmlElement("w:sz"); sz.set(qn("w:val"), "20"); rp.append(sz)
        f = OxmlElement("w:rFonts")
        for k in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            f.set(qn(k), YAZI)
        rp.append(f); nr.append(rp)
        t = OxmlElement("w:t"); t.text = "Videoyu aç"; nr.append(t); h.append(nr); p._p.append(h)

    adlar = list(dict.fromkeys(pp["ad"] for pp in paragraflar)) or list(veri["etiket"].values())
    p = par("", 10, sonra=0)
    font_ayarla(p.add_run("Konuşmacılar  "), 10, False, GRI)
    for i, ad in enumerate(adlar):
        font_ayarla(p.add_run(("   " if i else "") + "● "), 9, False, RENKLER[i % len(RENKLER)])
        font_ayarla(p.add_run(ad), 10, True)
    if veri.get("uyarilar"):
        par("Kontrol notları", 10, True, KOYU, once=10, sonra=2)
        for u in veri["uyarilar"]:
            par("– " + u, 10, False, GRI, sonra=1)

    def madde(isaret, metin_, ek=""):
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.left_indent, p.paragraph_format.first_line_indent = Cm(0.6), Cm(-0.6)
        font_ayarla(p.add_run(isaret), 11, False, GRI); font_ayarla(p.add_run(metin_), 11)
        if ek:
            font_ayarla(p.add_run("  " + ek), 9.5, False, GRI)
        return p

    # Toplantı: Kısaca (1-2 paragraf) + Kim ne yapacak (kisaca.py)
    ks = veri.get("kisaca")
    if ks:
        par("Kısaca", 13, True, KOYU, once=14, sonra=4)
        for pg in ks.get("ozet") or []:
            p = par(pg, 11, False, KOYU, sonra=6); p.paragraph_format.line_spacing = 1.3
        zk = lambda x: zaman_uzun(x["sn"], False) if x.get("sn") is not None else ""  # noqa: E731
        if ks.get("kararlar"):
            par(_t("Kararlar"), 11, True, KOYU, once=6, sonra=3)
            for x in ks["kararlar"]:
                madde("•  ", x["karar"], " · ".join(e for e in (x.get("musteri"), x.get("kim"), zk(x), *map(_zincir_metni, _zincirler(x))) if e))
        if ks.get("kisiler") or ks.get("belirsiz"):
            par("Kim ne yapacak", 11, True, KOYU, once=6, sonra=3)
            for k in (ks.get("kisiler") or []) + ([{"ad": "Sorumlu belirtilmedi", "isler": ks["belirsiz"]}] if ks.get("belirsiz") else []):
                par(k["ad"], 10.5, True, KOYU, once=4, sonra=1)
                for x in k["isler"]:
                    madde("☑  " if x.get("bitti") else "☐  ", x["is"], " · ".join(e for e in (x.get("musteri"), x.get("ne_zaman") or ("tarih belirtilmedi" if ks.get("surum", 1) >= 2 else ""), zk(x)) if e))
        if ks.get("oneriler"):
            par("Öneriler (karar değil)", 11, True, KOYU, once=6, sonra=3)
            for x in ks["oneriler"]:
                madde("–  ", x["oneri"], " · ".join(e for e in (x.get("musteri"), x.get("kim"), zk(x)) if e))
        if ks.get("acik_sorular"):
            par("Açık sorular", 11, True, KOYU, once=6, sonra=3)
            for x in ks["acik_sorular"]:
                madde("?  ", x["soru"], " · ".join(e for e in (x.get("musteri"), x.get("kime"), zk(x)) if e))

    # Özet
    o = veri.get("ozet")
    if o:
        par(_t("Özet"), 13, True, KOYU, once=14, sonra=4)
        p = par(o["ozet"], 11, False, KOYU, sonra=6); p.paragraph_format.line_spacing = 1.3
        if o.get("maddeler"):
            par(_t("Öne çıkanlar"), 11, True, KOYU, once=6, sonra=3)
            for m in o["maddeler"]:
                p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.left_indent, p.paragraph_format.first_line_indent = Cm(0.5), Cm(-0.5)
                font_ayarla(p.add_run("•  "), 11, False, GRI); font_ayarla(p.add_run(m), 11)
        for baslik_, anahtar, isaret in ((_t("Kararlar"), "kararlar", "•  "), (_t("Yapılacaklar"), "yapilacaklar", "☐  ")):
            if o.get(anahtar):
                par(baslik_, 11, True, KOYU, once=6, sonra=3)
                for m in o[anahtar]:
                    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(2)
                    p.paragraph_format.left_indent, p.paragraph_format.first_line_indent = Cm(0.5), Cm(-0.5)
                    font_ayarla(p.add_run(isaret), 11, False, GRI); font_ayarla(p.add_run(m), 11)
        if o.get("bolumler"):
            par(_t("Bölümler"), 11, True, KOYU, once=8, sonra=3)
            sb = any(b["sn"] >= 3600 for b in o["bolumler"])
            for b in o["bolumler"]:
                p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.left_indent, p.paragraph_format.first_line_indent = Cm(1.8 if sb else 1.4), Cm(-1.8 if sb else -1.4)
                p.paragraph_format.tab_stops.add_tab_stop(Cm(1.8 if sb else 1.4), WD_TAB_ALIGNMENT.LEFT)
                font_ayarla(p.add_run(zaman_uzun(b["sn"], sb) + "\t"), 9.5, False, GRI, tabular=True); font_ayarla(p.add_run(b["baslik"]), 11)

    # Kişisel: kariyer danışmanı notu (danisman.py) — kendi sayfasında
    dn = veri.get("danisman")
    if dn and yalniz_danisman:
        h = doc.add_heading(_t("Kişisel · Kariyer danışmanı notu"), level=1)
        for r in h.runs:
            font_ayarla(r, 16, True, KOYU)
        par("Yalnızca senin için. Bu dosya Drive'a gönderilmez.", 9.5, False, GRI, sonra=8)
        if dn.get("ozet"):
            p = par(dn["ozet"], 12.5, False, KOYU, sonra=8); p.paragraph_format.line_spacing = 1.3
        if dn.get("aksiyonlar"):
            par("Bu hafta yap", 11.5, True, KOYU, once=6, sonra=3)
            for a_ in dn["aksiyonlar"]:
                madde("☐  ", a_.get("ne", ""), a_.get("ne_zaman", ""))
        if dn.get("anlar"):
            par("Daha iyi olabilirdi", 11.5, True, KOYU, once=10, sonra=3)
            for a_ in dn["anlar"]:
                madde(f"{a_.get('zaman', '')}  ", a_.get("ne_oldu", ""))
                p = par("Şöyle deneyebilirdin: " + a_.get("daha_iyi", ""), 10.5, True, KOYU, sonra=1); p.paragraph_format.left_indent = Cm(1.2)
                p = par(a_.get("neden", ""), 10, False, GRI, sonra=6); p.paragraph_format.left_indent = Cm(1.2)
        if dn.get("guclu"):
            par("İyi giden", 11.5, True, KOYU, once=10, sonra=3)
            for g in dn["guclu"]:
                madde(f"{g.get('zaman', '')}  ", g.get("ne", ""))
        if dn.get("firsatlar"):
            par("Fırsatlar", 11.5, True, KOYU, once=10, sonra=3)
            for f_ in dn["firsatlar"]:
                par(f_.get("baslik", ""), 11, True, KOYU, once=3, sonra=1)
                par(f_.get("neden", ""), 10.5, False, KOYU, sonra=1)
                for ad_ in (f_.get("adimlar") or ([f_["nasil"]] if f_.get("nasil") else [])):
                    p = par("→ " + ad_, 10.5, False, GRI, sonra=1); p.paragraph_format.left_indent = Cm(0.6)
        if dn.get("dersler"):
            par("Ders", 11.5, True, KOYU, once=10, sonra=3)
            for d_ in dn["dersler"]:
                madde("•  ", f"{d_.get('baslik', '')}. {d_.get('ders', '')}")
        if dn.get("baglanti"):
            par("Geçen seferlerden: " + dn["baglanti"], 10, False, GRI, once=8, sonra=4)

    if veri.get("goruntu"):
        par(_t("Görüntü notları"), 13, True, KOYU, once=14, sonra=4)
        sg = any(g["sn"] >= 3600 for g in veri["goruntu"])
        for g in veri["goruntu"]:
            if not (g.get("aciklama") or g.get("yazi")):
                continue
            p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.left_indent, p.paragraph_format.first_line_indent = Cm(1.8 if sg else 1.4), Cm(-1.8 if sg else -1.4)
            p.paragraph_format.tab_stops.add_tab_stop(Cm(1.8 if sg else 1.4), WD_TAB_ALIGNMENT.LEFT)
            font_ayarla(p.add_run(zaman_uzun(g["sn"], sg) + "\t"), 9.5, False, GRI, tabular=True)
            if g.get("aciklama"):
                font_ayarla(p.add_run(g["aciklama"].strip() + ("  " if g.get("yazi") else "")), 10.5)
            if g.get("yazi"):
                font_ayarla(p.add_run("Ekranda: " + " / ".join(g["yazi"])), 10, False, GRI)

    # Ayırıcı çizgi
    p = par("", 6, once=6, sonra=12)
    pPr = p._p.get_or_add_pPr(); bdr = OxmlElement("w:pBdr"); bt = OxmlElement("w:bottom")
    bt.set(qn("w:val"), "single"); bt.set(qn("w:sz"), "4"); bt.set(qn("w:color"), "E8E8ED"); bdr.append(bt); pPr.append(bdr)

    # Konuşma: asılı girinti — zaman sol olukta, ad ve metin hizalı
    oluk = Cm(2.1 if saatli else 1.6)
    for pp in paragraflar:
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before, pf.space_after = Pt(0), Pt(10)
        pf.left_indent, pf.first_line_indent = oluk, -oluk
        pf.tab_stops.add_tab_stop(oluk, WD_TAB_ALIGNMENT.LEFT)
        pf.keep_together = True
        font_ayarla(p.add_run(zaman_uzun(pp["bas"], saatli) + "\t"), 9, False, GRI, tabular=True)
        renk = RENKLER[adlar.index(pp["ad"]) % len(RENKLER)] if pp["ad"] in adlar else KOYU
        font_ayarla(p.add_run(pp["ad"]), 11, True, renk)
        font_ayarla(p.add_run("\n\t" + pp["metin"]), 11)

    # Düz metin: ayrı sayfada, gezinti bölmesinde görünen başlıkla
    h = doc.add_heading(_t("Düz metin"), level=1)
    h.paragraph_format.page_break_before = True
    for r in h.runs:
        font_ayarla(r, 16, True, KOYU)
    p = par(" ".join(pp["metin"] for pp in paragraflar), 10.5, False, KOYU, once=6)
    p.paragraph_format.line_spacing = 1.3

    # Alt bilgi: başlık + sayfa numarası
    fp = doc.sections[0].footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font_ayarla(fp.add_run((veri["baslik"] or "Transkript")[:70] + "   ·   "), 8, False, GRI)
    r = fp.add_run(); font_ayarla(r, 8, False, GRI, tabular=True)
    _alan_ekle(r, "PAGE")
    doc.save(yol)


def _srt_zaman(sn):
    sn = max(0.0, float(sn)); h = int(sn // 3600); m = int(sn % 3600 // 60); s = sn % 60
    return f"{h:02d}:{m:02d}:{int(s):02d},{int(round((s - int(s)) * 1000)):03d}"


SATIR_EN_COK = 42
BAGLAC = {"ve", "ama", "fakat", "çünkü", "ki", "ya", "de", "da", "veya", "ancak", "and", "but", "or", "because", "so"}


def _satira_bol(metin):
    """Metni en çok iki satıra, 42 karakteri aşmadan, doğal bir yerden böler."""
    if len(metin) <= SATIR_EN_COK:
        return [metin]
    kelimeler = metin.split()
    en_iyi, en_iyi_puan = None, None
    for i in range(1, len(kelimeler)):
        a, b = " ".join(kelimeler[:i]), " ".join(kelimeler[i:])
        puan = abs(len(a) - len(b)) + (0 if max(len(a), len(b)) <= SATIR_EN_COK else 100 + 10 * (max(len(a), len(b)) - SATIR_EN_COK))
        if a[-1] in ",;:.!?…":
            puan -= 8
        if kelimeler[i].lower().strip(",.") in BAGLAC:
            puan -= 5
        if en_iyi_puan is None or puan < en_iyi_puan:
            en_iyi, en_iyi_puan = (a, b), puan
    return list(en_iyi) if en_iyi else [metin]


def altyazi_parcalari(segmentler):
    """YouTube'a uygun altyazı parçaları: konuşmacı adı yok, 1-2 satır, satır ≤42 karakter, 1-7 sn.
    Kelime zamanları varsa onları, yoksa karakter oranıyla yaklaşık zamanları kullanır."""
    kelimeler = []
    for s in segmentler:
        if s.get("kelimeler"):
            kelimeler += [(k[0], k[1], k[2]) for k in s["kelimeler"]]
        else:
            parca = s["metin"].split()
            if not parca:
                continue
            toplam = sum(len(p) + 1 for p in parca); t = s["bas"]; sure = max(0.5, s["bit"] - s["bas"])
            for p in parca:
                dt = sure * (len(p) + 1) / toplam
                kelimeler.append((t, t + dt, p)); t += dt
    parcalar, simdiki = [], []

    def kapat():
        if simdiki:
            parcalar.append({"bas": simdiki[0][0], "bit": simdiki[-1][1], "metin": " ".join(k[2] for k in simdiki),
                             "kelimeler": [(k[0], k[1], k[2]) for k in simdiki]})
            simdiki.clear()

    for i, (b, e, w) in enumerate(kelimeler):
        if simdiki:
            metin = " ".join(k[2] for k in simdiki)
            sure = e - simdiki[0][0]
            bosluk = b - simdiki[-1][1]
            onceki = simdiki[-1][2]
            if (len(metin) + 1 + len(w) > 2 * SATIR_EN_COK - 6 or sure > 6.5 or bosluk > 1.0 and len(metin) >= 15
                    or onceki[-1:] in ".?!…" and len(metin) >= 20 and (e - simdiki[0][0]) >= 1.0
                    or onceki[-1:] in ",;:" and len(metin) >= 60):
                kapat()
        simdiki.append((b, e, w))
    kapat()
    # Çok kısa parçaları (1 sn altı) sığıyorsa sonrakiyle birleştir
    birlesik = []
    for p in parcalar:
        if birlesik and (birlesik[-1]["bit"] - birlesik[-1]["bas"]) < 1.0 and len(birlesik[-1]["metin"]) + 1 + len(p["metin"]) <= 2 * SATIR_EN_COK - 6 \
                and p["bas"] - birlesik[-1]["bit"] < 1.5:
            birlesik[-1]["metin"] += " " + p["metin"]; birlesik[-1]["bit"] = p["bit"]
        else:
            birlesik.append(p)
    parcalar = birlesik
    # Zamanları düzelt: en az 1,2 sn, bir sonrakiyle çakışma yok
    for i, p in enumerate(parcalar):
        p["bit"] = max(p["bit"] + 0.15, p["bas"] + 1.2)
        if i + 1 < len(parcalar):
            p["bit"] = min(p["bit"], parcalar[i + 1]["bas"] - 0.04)
            if p["bit"] <= p["bas"]:
                p["bit"] = p["bas"] + 0.6
    return parcalar


def yaz_srt(yol, veri, segmentler):
    """YouTube'a doğrudan yüklenebilir altyazı dosyası (.srt). Anlamlı kırılım planı (altyazi.py) varsa satırlar
    anlam birimlerinden bölünür; yoksa kural tabanlı bölme."""
    parcalar = None
    try:
        import altyazi
        plan = altyazi.gecerli_plan(veri or {}, segmentler)
        if plan:
            parcalar = altyazi.parcalar(segmentler, plan)
    except Exception:  # noqa: BLE001
        parcalar = None
    with open(yol, "w", encoding="utf-8") as f:
        for n, p in enumerate(parcalar or altyazi_parcalari(segmentler), 1):
            f.write(f"{n}\n{_srt_zaman(p['bas'])} --> {_srt_zaman(p['bit'])}\n" + "\n".join(p.get("satirlar") or _satira_bol(p["metin"])) + "\n\n")


def senaryo_sahneleri(veri, paragraflar, segmentler=None):
    """Sahne listesi: [{"n","sn","goruntu","super":[...],"ses":[(ad, metin)]}]
    Konuşma, cümle düzeyinde (segmentler) sahnelere dağıtılır; aynı konuşmacının ardışık cümleleri birleşir."""
    kareler = sorted(veri.get("goruntu") or [], key=lambda g: g["sn"])
    if not kareler:
        kareler = [{"sn": 0, "yazi": [], "aciklama": ""}]
    etiket = veri.get("etiket", {})
    if segmentler:
        cumleler = [(s["bas"], etiket.get(str(s["k"]), "Konuşmacı"), s["metin"]) for s in segmentler]
    else:
        cumleler = [(p["bas"], p["ad"], p["metin"]) for p in paragraflar]
    sahneler = []
    for i, k in enumerate(kareler):
        bas = k["sn"]; bit = kareler[i + 1]["sn"] if i + 1 < len(kareler) else 10 ** 9
        ses = []
        for cb, ad, metin in cumleler:
            if (bas - 0.3 <= cb < bit - 0.3) if i else (cb < bit - 0.3):
                if ses and ses[-1][0] == ad:
                    ses[-1] = (ad, ses[-1][1] + " " + metin)
                else:
                    ses.append((ad, metin))
        goruntu = (k.get("aciklama") or "").strip()
        yazi = list(k.get("yazi") or [])
        # Betim, ekrandaki yazıyı tekrar ediyorsa (SÜPER sütununda zaten var) kısalt
        if yazi and goruntu:
            kelimeler = {w.lower() for y in yazi for w in re.findall(r"\w+", y) if len(w) > 2}
            betim_kelime = [w.lower() for w in re.findall(r"\w+", goruntu) if len(w) > 2]
            if betim_kelime and sum(1 for w in betim_kelime if w in kelimeler) / len(betim_kelime) >= 0.4:
                goruntu = "Tam ekran yazı kartı; metin SÜPER sütununda."
        sahneler.append({"n": i + 1, "sn": bas, "goruntu": goruntu, "super": yazi, "ses": ses})
    return sahneler


def yaz_senaryo_docx(yol, veri, paragraflar, segmentler=None):
    """Prodüksiyon senaryosu: sahne sahne GÖRÜNTÜ / SÜPER / DIŞ SES (ya da DİYALOG) sütunları."""
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    GRI = RGBColor(0x6E, 0x6E, 0x73); KOYU = RGBColor(0x1D, 0x1D, 0x1F); YAZI = "Helvetica Neue"
    sahneler = senaryo_sahneleri(veri, paragraflar, segmentler)
    konusmacilar = list(dict.fromkeys(p["ad"] for p in paragraflar))
    diyalog = len(konusmacilar) > 1
    ses_baslik = "DİYALOG" if diyalog else "DIŞ SES"
    saatli = any(sh["sn"] >= 3600 for sh in sahneler)

    doc = Document()
    for sec in doc.sections:
        sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)        # yatay A4: üç sütun rahat sığsın
        sec.orientation = 1
        sec.top_margin = sec.bottom_margin = Cm(1.8); sec.left_margin = sec.right_margin = Cm(2.0)
    normal = doc.styles["Normal"]; normal.font.name = YAZI; normal.font.size = Pt(10.5); normal.font.color.rgb = KOYU
    normal.paragraph_format.space_after = Pt(0); normal.paragraph_format.line_spacing = 1.2
    rpr = normal.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.append(rf)
    for k in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(k), YAZI)
    lang = OxmlElement("w:lang"); lang.set(qn("w:val"), "tr-TR"); rpr.append(lang)
    cp = doc.core_properties; cp.title = (veri.get("baslik") or "Senaryo"); cp.author = "Transkript"; cp.language = "tr-TR"

    def run(p, metin, boyut=10.5, kalin=False, italik=False, renk=None):
        r = p.add_run(metin); r.font.size = Pt(boyut); r.bold = kalin; r.italic = italik; r.font.name = YAZI
        if renk is not None:
            r.font.color.rgb = renk
        return r

    p = doc.add_paragraph(); run(p, "SENARYO", 9, True, renk=GRI)
    sp = OxmlElement("w:spacing"); sp.set(qn("w:val"), "30"); p.runs[0]._r.get_or_add_rPr().append(sp)
    p = doc.add_paragraph(); run(p, veri.get("baslik") or "Senaryo", 18, True); p.paragraph_format.space_after = Pt(2)
    bilgi = [veri.get("kanal"), sure_metni(veri["sure"]) if veri.get("sure") else None,
             f"{len(sahneler)} sahne", ("Konuşmacılar: " + ", ".join(konusmacilar)) if konusmacilar else None]
    p = doc.add_paragraph(); run(p, "  ·  ".join(b for b in bilgi if b), 9.5, renk=GRI); p.paragraph_format.space_after = Pt(10)

    tablo = doc.add_table(rows=1, cols=4)
    tablo.style = "Table Grid"; tablo.alignment = WD_TABLE_ALIGNMENT.CENTER
    genislik = [Cm(2.4), Cm(8.6), Cm(5.6), Cm(9.1)]
    basliklar = ["SAHNE", "GÖRÜNTÜ", "SÜPER", ses_baslik]
    for i, (h, g) in enumerate(zip(basliklar, genislik)):
        c = tablo.rows[0].cells[i]; c.width = g
        c.paragraphs[0].text = ""; run(c.paragraphs[0], h, 9, True, renk=GRI)
        tcPr = c._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), "F2F2F4"); tcPr.append(shd)

    def hucre_yaz(c, parcalar):
        """parcalar: [(metin, boyut, kalin, italik, renk)]; her biri ayrı paragraf."""
        c.paragraphs[0].text = ""
        ilk = True
        for metin, boyut, kalin, italik, renk in parcalar:
            p = c.paragraphs[0] if ilk else c.add_paragraph(); ilk = False
            p.paragraph_format.space_after = Pt(3)
            run(p, metin, boyut, kalin, italik, renk)
        if ilk:
            run(c.paragraphs[0], "—", 10.5, renk=GRI)

    for sh in sahneler:
        satir = tablo.add_row(); hucreler = satir.cells
        for i, g in enumerate(genislik):
            hucreler[i].width = g
        hucre_yaz(hucreler[0], [(f"{sh['n']}", 11, True, False, None), (zaman_uzun(sh["sn"], saatli), 9, False, False, GRI)])
        hucre_yaz(hucreler[1], [(sh["goruntu"], 10.5, False, False, None)] if sh["goruntu"] else [])
        hucre_yaz(hucreler[2], [(y, 10.5, False, False, None) for y in sh["super"]])
        ses = []
        for ad, metin in sh["ses"]:
            if diyalog:
                ses.append((__import__('harfler').ust(ad), 9, True, False, GRI))
            ses.append((metin, 10.5, False, False, None))
        hucre_yaz(hucreler[3], ses)
        # satır sayfa sonunda bölünmesin
        trPr = satir._tr.get_or_add_trPr(); cs = OxmlElement("w:cantSplit"); trPr.append(cs)

    fp = doc.sections[0].footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run((veri.get("baslik") or "Senaryo")[:70] + "   ·   "); r.font.size = Pt(8); r.font.color.rgb = GRI
    r = fp.add_run(); r.font.size = Pt(8); r.font.color.rgb = GRI; _alan_ekle(r, "PAGE")
    doc.save(yol)


def yaz_senaryo_txt(yol, veri, paragraflar, segmentler=None):
    """Aynı senaryo düz metin olarak (kopyalamak / sohbette göstermek için)."""
    sahneler = senaryo_sahneleri(veri, paragraflar, segmentler)
    diyalog = len({p["ad"] for p in paragraflar}) > 1
    saatli = any(sh["sn"] >= 3600 for sh in sahneler)
    with open(yol, "w", encoding="utf-8") as f:
        f.write(f"SENARYO — {veri.get('baslik') or ''}\n\n")
        for sh in sahneler:
            f.write(f"SAHNE {sh['n']}  ·  {zaman_uzun(sh['sn'], saatli)}\n")
            f.write("  GÖRÜNTÜ: " + (sh["goruntu"] or "—") + "\n")
            f.write("  SÜPER: " + (" / ".join(sh["super"]) if sh["super"] else "—") + "\n")
            f.write("  " + ("DİYALOG" if diyalog else "DIŞ SES") + ":\n")
            if sh["ses"]:
                for ad, metin in sh["ses"]:
                    f.write(("    " + __import__('harfler').ust(ad) + ": " if diyalog else "    ") + metin + "\n")
            else:
                f.write("    —\n")
            f.write("\n")


ERTELE_SN = 0.0            # site süreci 2.0 yapar: docx/srt (0,6 sn) ad/paragraf değişikliğinde hemen değil, 2 sn sonra tek seferde üretilir
_ertelenen = {}            # temel → threading.Timer
_ertele_kilit = threading.Lock()


def _belgeleri_yaz(temel_yol, veri, segmentler, paragraflar):
    yaz_docx(str(temel_yol) + ".docx", veri, paragraflar)
    if veri.get("danisman"):                      # kişisel danışman notu AYRI dosyada (Drive'a gitmez; 2026-10-23)
        yaz_docx(str(temel_yol) + ".danisman.docx", veri, [], yalniz_danisman=True)
        try:
            import danisman as _dn
            Path(str(temel_yol) + ".danisman.txt").write_text(_dn.metin(veri["danisman"]) + "\n", encoding="utf-8")
        except Exception:  # noqa: BLE001
            pass
    yaz_srt(str(temel_yol) + ".srt", veri, segmentler)
    if veri.get("goruntu"):
        yaz_senaryo_docx(str(temel_yol) + ".senaryo.docx", veri, paragraflar, segmentler)
        yaz_senaryo_txt(str(temel_yol) + ".senaryo.txt", veri, paragraflar, segmentler)


def bekleyenleri_yaz(temel=None):
    """Ertelenmiş docx/srt üretimini şimdi yapar (indirme/zip/dışa aktarma öncesi). temel=None: hepsi."""
    with _ertele_kilit:
        secilen = [k for k in _ertelenen if temel is None or k == str(temel)]
        isler = [(k, _ertelenen.pop(k)) for k in secilen]
    for _k, (t, arg) in isler:
        t.cancel()
        _belgeleri_yaz(*arg)


def ciktilari_yaz(temel_yol, veri, segmentler):
    """txt + json hemen; docx + srt (+senaryo) hemen ya da ERTELE_SN sonra (toplantılar hep hemen: Drive kopyası docx'i bekler).
    isimlendir.py ve site de bunu kullanır."""
    paragraflar = paragrafla(segmentler, veri["etiket"])
    yaz_txt(str(temel_yol) + ".txt", veri, paragraflar)
    import kayit_json                                        # 2026-10-25: süreçler arası kilit + geçici dosya + girintisiz JSON
    with kayit_json.kilit(str(temel_yol) + ".json"):
        kayit_json.yaz(str(temel_yol) + ".json", {"veri": veri, "segmentler": segmentler})
    if ERTELE_SN > 0 and not veri.get("toplanti"):
        anahtar = str(temel_yol)
        with _ertele_kilit:
            eski = _ertelenen.pop(anahtar, None)
            if eski:
                eski[0].cancel()
            t = threading.Timer(ERTELE_SN, lambda: (bekleyenleri_yaz(anahtar)))
            t.daemon = True
            _ertelenen[anahtar] = (t, (temel_yol, veri, segmentler, paragraflar))
            t.start()
    else:
        _belgeleri_yaz(temel_yol, veri, segmentler, paragraflar)
    return paragraflar


# ----------------------------------------------------------------------------- ana akış
def main():
    args = [a for a in sys.argv[1:] if a.strip()]
    if not args:
        print(__doc__)
        raise SystemExit(1)

    isimler, konusmaci_sayisi, altyazi_iste, ad = [], None, True, None
    ozet_iste, ozet_dili, konusmaci_ayir, goruntu = True, "tr", True, False
    pozisyonel = []
    i = 0
    while i < len(args):
        a = args[i]
        if a == "--isimler":
            isimler = [x.strip() for x in args[i + 1].split(",") if x.strip()]; i += 2
        elif a == "--konusmaci":
            konusmaci_sayisi = int(args[i + 1]); i += 2
        elif a == "--konusmaci-yok":
            konusmaci_ayir = False; i += 1
        elif a == "--altyazi-yok":
            altyazi_iste = False; i += 1
        elif a == "--ad":
            ad = args[i + 1]; i += 2
        elif a == "--cerez":
            global COOKIES_BROWSER
            COOKIES_BROWSER = args[i + 1].lower(); i += 2
        elif a == "--ozet-yok":
            ozet_iste = False; i += 1
        elif a == "--claude-yok":                         # Meet örnek toplantı testi (2026-10-25): Claude çağrısı yok, ad önerisi ve altyazı planı da yok
            os.environ["VH_CLAUDE_YOK"] = "1"; i += 1
        elif a == "--goruntu":
            goruntu = True; i += 1
        elif a == "--ozet-dili":
            ozet_dili = args[i + 1].lower(); i += 2
        elif a == "--ses-temizle":
            global SES_TEMIZLE
            SES_TEMIZLE = True; i += 1
        elif a == "--sozluk":
            global EK_TERIMLER
            EK_TERIMLER = [x.strip() for x in args[i + 1].split(",") if x.strip()]; i += 2
        elif a in ("--terim-yok", "--terim-duzelt"):
            global TERIM_DUZELT
            TERIM_DUZELT = a == "--terim-duzelt"; i += 1
        else:
            pozisyonel.append(a); i += 1

    link = pozisyonel[0]
    dil = pozisyonel[1].lower() if len(pozisyonel) > 1 else "auto"
    model_kisa = pozisyonel[2].lower() if len(pozisyonel) > 2 else "turbo"
    if dil not in ("auto", "multi") and not re.match(r"^[a-z]{2,3}$", dil):
        raise SystemExit("HATA: dil kodu geçersiz (örn. tr, en, de, fr, es, auto, multi).")
    if model_kisa not in MODELS:
        raise SystemExit("HATA: model sadece turbo veya best olabilir.")
    if konusmaci_sayisi is None and isimler:
        konusmaci_sayisi = len(isimler)

    # Kontrol noktası (2026-10-19): indirme ve ses çözümü (+ konuşmacı ayırma) sonuçları işe özel klasöre yazılır; site iş ortasında
    # yeniden başlarsa yeniden sıraya giren iş bu adımları atlar. İş başarıyla bitince klasör silinir (eskiler site döngüsünde 3 günde).
    import hashlib
    # yerel dosyada boyut ve zaman da anahtara girer (2026-10-27): birleştirilen toplantı kaydı ilk parçayla aynı yolu alıyor, eski parçanın sesi kullanılıyordu
    _yerel = os.stat(link) if os.path.isfile(link) else None
    _ek = f"|{_yerel.st_size}|{int(_yerel.st_mtime)}" if _yerel else ""
    kn = KLASOR / ".onbellek" / "is" / hashlib.sha1(f"{link}|{dil}|{model_kisa}|{konusmaci_ayir}|{konusmaci_sayisi}|{SES_TEMIZLE}{_ek}".encode()).hexdigest()[:16]
    kn_ses, kn_bilgi, kn_cozum = kn / "ses.m4a", kn / "bilgi.json", kn / "cozum.json"
    with tempfile.TemporaryDirectory(prefix="transkript_") as gecici:
        if kn_ses.exists() and kn_bilgi.exists() and not goruntu:
            ses, altyazi, bilgi = kn_ses, None, json.loads(kn_bilgi.read_text(encoding="utf-8"))
            log("[1/5] Önceki denemeden kalan ses kullanılıyor (yeniden indirilmedi).")
        else:
            ses, altyazi, bilgi = indir(link, gecici, dil, altyazi_iste, video_iste=goruntu)
            if not altyazi and not goruntu:
                try:
                    kn.mkdir(parents=True, exist_ok=True)
                    shutil.copy2(ses, kn_ses)
                    kn_bilgi.write_text(json.dumps({k: x for k, x in bilgi.items() if k != "video"}, ensure_ascii=False, default=str), encoding="utf-8")
                except Exception:  # noqa: BLE001
                    pass
        if SES_TEMIZLE and not altyazi:
            try:
                import temizle as temizle_modulu
                log("      Ses temizleniyor (gürültü bastırma)...")
                ses = temizle_modulu.temizle(ses, Path(gecici) / "ses_temiz.m4a", "orta", log=log)
            except Exception as e:  # noqa: BLE001
                log(f"      Ses temizlenemedi ({e}); olduğu gibi sürülüyor")
        log(f"      Başlık: {bilgi['baslik']}")
        log(f"      Süre: {int(bilgi['sure'])} sn")
        ayirici = None
        onceki_cozum = None
        if kn_cozum.exists() and not altyazi:
            try:
                onceki_cozum = json.loads(kn_cozum.read_text(encoding="utf-8"))
            except Exception:  # noqa: BLE001
                onceki_cozum = None
        if onceki_cozum:
            log("[2/5] Önceki denemeden kalan konuşmacı ayrımı kullanılıyor.")
            log("[3/5] Önceki denemeden kalan yazıya dökme kullanılıyor (yeniden çözülmedi).")
            parcalar = onceki_cozum["parcalar"]
            konusmacilari_ayir.izler = onceki_cozum.get("izler") or {}
        elif konusmaci_ayir and not altyazi and motor_sec() == "mlx":
            # Ayırma (CPU) ile yazıya dökme (GPU) aynı anda: ölçüm 17 dk ses → ayırma süresi tamamen gizleniyor
            ayirici = _ayirma_baslat(ses, konusmaci_sayisi, Path(gecici) / "ayirma.json")
            parcalar = []
        elif konusmaci_ayir:
            parcalar = konusmacilari_ayir(ses, konusmaci_sayisi)
        else:
            log("[2/5] Konuşmacı ayırma atlandı.")
            parcalar = []
        if onceki_cozum:
            tespit_dil, metinler, kaynak, diller = (onceki_cozum["tespit_dil"], onceki_cozum["metinler"], onceki_cozum["kaynak"],
                                                    onceki_cozum["diller"])
        elif altyazi:
            log(f"[3/5] Videoda yüklü altyazı var ({altyazi.name}), sesten çözmeye gerek kalmadı.")
            metinler = altyazi_oku(altyazi)
            tespit_dil = altyazi_dili(altyazi)
            kaynak = "videodaki yüklü altyazı"
            diller = [tespit_dil]
        else:
            tespit_dil, metinler = transkript_et(ses, dil, MODELS[model_kisa], isimler, parcalar)
            kaynak = f"sesten çözüm ({MODELS[model_kisa]})"
            diller = getattr(transkript_et, "diller", [tespit_dil])
        if ayirici is not None:
            parcalar = _ayirma_bitir(ayirici, ses, konusmaci_sayisi)
        if not onceki_cozum and not altyazi and kn.exists():
            try:
                kn_cozum.write_text(json.dumps({"tespit_dil": tespit_dil, "metinler": metinler, "kaynak": kaynak, "diller": diller,
                                                "parcalar": parcalar, "izler": getattr(konusmacilari_ayir, "izler", {}) or {}},
                                               ensure_ascii=False, default=lambda o: o.tolist() if hasattr(o, "tolist") else float(o)),
                                    encoding="utf-8")
            except Exception as e:  # noqa: BLE001
                log(f"      (kontrol noktası yazılamadı: {str(e)[:80]})")
        goruntu_notlari, kare_klasoru_gecici = [], None
        if goruntu:
            if bilgi.get("video"):
                import goruntu as goruntu_modulu
                log("[3/5] Görüntü inceleniyor...")
                kare_klasoru_gecici = Path(gecici) / "kareler"
                goruntu_notlari = goruntu_modulu.incele(bilgi["video"], kare_klasoru_gecici, bilgi.get("sure") or 0,
                                                        tespit_dil if tespit_dil in ("tr", "en") else "tr", log=log)
            else:
                log("      Görüntü için video alınamadı; görüntü incelemesi atlandı.")
        kareler_tasinacak = sorted(kare_klasoru_gecici.glob("k_*.jpg")) if kare_klasoru_gecici else []
        kare_yedek = {p.name: p.read_bytes() for p in kareler_tasinacak}

    log("[4/5] Konuşmacılar eşleniyor, dosyalar yazılıyor...")
    segmentler = konusmaci_eslestir(metinler, parcalar)
    for s in segmentler:
        s["k"] = str(s["k"])
    if not konusmaci_sayisi and konusmaci_ayir:
        segmentler = hayalet_temizle(segmentler)
    if TERIM_DUZELT and tespit_dil == "tr" and not altyazi and segmentler:
        try:
            import terimler as terimler_modulu
            if terimler_modulu.ozet_hazir():
                log("      İngilizce terimler denetleniyor...")
                yapilan = terimler_modulu.duzelt(segmentler, log=log, ekstra=EK_TERIMLER, dil=tespit_dil)
                if yapilan:
                    log("      Düzeltilen terimler: " + ", ".join(f"{a} → {b}" for a, b in yapilan[:12])
                        + (f" (+{len(yapilan) - 12})" if len(yapilan) > 12 else ""))
        except Exception as e:  # noqa: BLE001
            log(f"      Terim denetimi yapılamadı ({e}).")
    temizlik = []
    if not altyazi and segmentler:
        try:
            import claude_yaz
            if "/kayitlar/" in str(link):
                os.environ["VH_TOPLANTI"] = "1"                    # toplantı kaydı: Ayarlar'daki toplantı iznine uyulur
            if claude_yaz.hazir():
                import terimler as terimler_modulu
                log("      Metin Claude ile denetleniyor (yanlış duyulan adlar ve terimler)...")
                temizlik = terimler_modulu.claude_temizle(segmentler, dil=tespit_dil, isimler=list(isimler or []) + EK_TERIMLER, log=log)
                if temizlik:
                    log("      Düzeltilenler: " + ", ".join(f"{a} → {b}" for a, b in temizlik[:15])
                        + (f" (+{len(temizlik) - 15})" if len(temizlik) > 15 else ""))
                else:
                    log("      Düzeltilecek bir şey bulunmadı.")
        except Exception as e:  # noqa: BLE001
            log(f"      Metin denetimi atlandı ({str(e)[:100]}).")
    etiket = isimleri_esle(segmentler, isimler)
    izler = {str(k): v for k, v in (getattr(konusmacilari_ayir, "izler", {}) or {}).items()}
    sesler_v = {str(k): izler[str(k)] for k in etiket if str(k) in izler}
    taninan = []
    if sesler_v and not os.environ.get("VH_MISAFIR"):        # misafir işinde sahibin tanınan sesleri (kişileri) kullanılmaz
        import sesler as sesler_modulu
        for k, tanin_ad in sesler_modulu.tani(sesler_v, toplanti="/kayitlar/" in str(link)).items():   # toplantıda video sesleri kullanılmaz
            anahtar = k if k in etiket else (int(k) if str(k).isdigit() and int(k) in etiket else None)
            if anahtar is not None and str(etiket[anahtar]).startswith("Konuşmacı ") and tanin_ad not in etiket.values():
                etiket[anahtar] = tanin_ad
                taninan.append(tanin_ad)
        if taninan:
            log("      Sesinden tanındı: " + ", ".join(taninan))
    # Metin denetiminden sonra altyazı kırılımı ve özet birbirinden bağımsız: ad önerisiyle AYNI ANDA yürür (2026-10-19; eskiden
    # sırayla ~3-5 dk ekliyordu). Sonuçlar aşağıda veri'ye yazılır.
    from concurrent.futures import ThreadPoolExecutor as _Havuz
    _havuz = _Havuz(max_workers=2)

    def _altyazi_plani():
        import altyazi
        return altyazi.plan_uret(segmentler, tespit_dil, log=log)

    def _ozet():
        import ozet as ozet_modulu
        if not ozet_modulu.hazir():
            log("[5/5] Özet modeli hazır değil, özet atlandı.")
            return None
        log("[5/5] Özet çıkarılıyor...")
        return ozet_modulu.ozet_uret(paragrafla(segmentler, etiket), tespit_dil, ozet_dili, log=log, goruntu=goruntu_notlari or None)
    f_altyazi = _havuz.submit(_altyazi_plani) if segmentler and "/kayitlar/" not in str(link) and not os.environ.get("VH_CLAUDE_YOK") else None
    f_ozet = _havuz.submit(_ozet) if ozet_iste and segmentler else None
    ad_onerileri = []
    if any(str(v).startswith("Konuşmacı ") for v in etiket.values()) and len(etiket) >= 2 and not os.environ.get("VH_CLAUDE_YOK"):
        try:
            import adlar as adlar_modulu
            ad_onerileri = adlar_modulu.oner(segmentler, {str(k): v for k, v in etiket.items()}, log=log)
            if ad_onerileri:
                log("      Ad önerileri: " + ", ".join(f"{o['etiket']} → {o['ad']}" for o in ad_onerileri))
        except Exception as e:  # noqa: BLE001
            log(f"      Ad çıkarımı atlandı ({e}).")
    uyarilar = eksik_kontrol(segmentler, parcalar)
    if not segmentler:
        uyarilar.append("Videoda konuşma tespit edilemedi (müzik ya da sessiz olabilir).")
    veri = {
        "link": link, "baslik": bilgi["baslik"], "kanal": bilgi["kanal"], "sure": bilgi["sure"],
        "kapak": bilgi["kapak"], "platform": bilgi["platform"], "video_id": bilgi["video_id"],
        "dil": tespit_dil, "diller": diller, "kaynak": kaynak, "model": MODELS[model_kisa], "olusturma": time.time(), "surum": _surum(),
        "etiket": {str(k): v for k, v in etiket.items()},
        "uyarilar": uyarilar, "sesler": sesler_v, "taninan": taninan, "ad_onerileri": ad_onerileri,
    }
    if os.environ.get("VH_MISAFIR"):                          # Studio davetlisinin kaydı: yalnız o görür; sahibin listelerine ve Ghost'a girmez
        veri["misafir"] = os.environ["VH_MISAFIR"]
    if temizlik:
        veri["temizlik"] = {"duzeltme": len(temizlik), "ornekler": [f"{a} → {b}" for a, b in temizlik[:20]], "zaman": time.time()}
    if f_altyazi:                                             # videolarda anlamlı altyazı kırılımı (toplantıda gerek yok)
        try:
            plan = f_altyazi.result()
            if plan:
                veri["altyazi_kirilim"] = plan
                log("      Altyazı satırları anlam birimlerine göre bölündü.")
        except Exception as e:  # noqa: BLE001
            log(f"      Altyazı kırılımı atlandı ({str(e)[:80]}).")
    if isimler and sesler_v and not os.environ.get("VH_MISAFIR"):
        try:
            import sesler as sesler_modulu
            sesler_modulu.ogren(veri, log)
        except Exception:  # noqa: BLE001
            pass
    if not ad:
        kisa = re.sub(r"[^0-9A-Za-zğüşıöçĞÜŞİÖÇ ]+", " ", bilgi["baslik"] or "")
        kisa = re.sub(r"^[0-9a-f]{6} ", "", kisa)
        kisa = "_".join(kisa.split()[:5])
        ad = f"transkript_{tespit_dil}" + (f"_{kisa}" if kisa else "")
    temel = cikti_adi(ad)
    if goruntu_notlari:
        kk = KLASOR / "kareler" / temel.name
        kk.mkdir(parents=True, exist_ok=True)
        for adk, veri_b in kare_yedek.items():
            (kk / adk).write_bytes(veri_b)
        veri["goruntu"] = goruntu_notlari
    ciktilari_yaz(temel, veri, segmentler)

    if f_ozet:
        try:
            o = f_ozet.result()
            if o:
                veri["ozet"] = o
                ciktilari_yaz(temel, veri, segmentler)
            elif o is not None:
                log("      Özet üretilemedi; transkript özetsiz kaydedildi.")
        except Exception as e:  # noqa: BLE001
            log(f"      Özet çıkarılamadı ({e}); transkript özetsiz kaydedildi.")
    _havuz.shutdown(wait=False)

    if uyarilar:
        log("\nKontrol notları:")
        for u in uyarilar:
            log("  - " + u)
    shutil.rmtree(kn, ignore_errors=True)                       # iş bitti: kontrol noktası gereksiz
    log(f"\nBitti. Dosya: {temel}.txt")
    log(f"Word:  {temel}.docx")
    log(f"Konuşmacılar: " + ", ".join(veri["etiket"].values()))
    if not isimler:
        log("İsim vermek için: isimlendir.py " + Path(str(temel) + ".json").name
            + ' "Konuşmacı 1=Ad" "Konuşmacı 2=Ad"')


if __name__ == "__main__":
    main()
    # Grafik işlemci motoru ile görüntü kütüphanesi kapanışta çakışabiliyor (temizlik anında çökme);
    # iş bitmiş olduğu için çıktıyı boşaltıp temiz çıkış koduyla hemen kapan.
    sys.stdout.flush(); sys.stderr.flush()
    os._exit(0)
